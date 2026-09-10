"""PPTX 结构读取。"""

from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any, Dict, List

from ..protocol import OfficeError

try:
    from pptx import Presentation  # type: ignore

    PPTX_AVAILABLE = True
except Exception:  # pragma: no cover
    PPTX_AVAILABLE = False


def require_pptx() -> None:
    if not PPTX_AVAILABLE:
        raise OfficeError(
            "E_ENV_MISSING",
            "PPTX 处理需要 python-pptx（Tier1）",
            detail="python-pptx is not importable",
            remedy="调用 office_env_setup 安装 Tier1 组件，或先安装 python-pptx",
        )


def _open(path: Path):
    try:
        return Presentation(str(path))
    except zipfile.BadZipFile as exc:
        raise OfficeError(
            "E_DOC_CORRUPT", "PPTX 不是有效的 ZIP 容器", detail=str(exc)
        ) from exc
    except Exception as exc:
        raise OfficeError(
            "E_DOC_CORRUPT",
            "PPTX 打开失败",
            detail="%s: %s" % (type(exc).__name__, exc),
        ) from exc


def pptx_outline(path: Path, *, max_slides: int = 0) -> Dict[str, Any]:
    require_pptx()
    presentation = _open(path)
    slides: List[Dict[str, Any]] = []
    for slide_index, slide in enumerate(presentation.slides):
        shapes: List[Dict[str, Any]] = []
        for shape_index, shape in enumerate(slide.shapes):
            entry: Dict[str, Any] = {
                "index": shape_index,
                "shape_id": shape.shape_id,
                "name": shape.shape_name if hasattr(shape, "shape_name") else shape.name,
                "type": str(shape.shape_type),
                "has_text_frame": bool(shape.has_text_frame),
                "left_emu": shape.left,
                "top_emu": shape.top,
                "width_emu": shape.width,
                "height_emu": shape.height,
                "bounds_cm": {"left": round(shape.left / 360000, 3), "top": round(shape.top / 360000, 3),
                              "width": round(shape.width / 360000, 3), "height": round(shape.height / 360000, 3)},
                "out_of_bounds": (shape.left < 0 or shape.top < 0 or shape.left + shape.width > presentation.slide_width
                                  or shape.top + shape.height > presentation.slide_height),
            }
            if shape.has_text_frame:
                entry["text"] = shape.text_frame.text
                entry["paragraphs"] = [{"text": p.text, "level": p.level,
                    "runs": [{"text": r.text, "font_name": r.font.name,
                              "size_pt": r.font.size.pt if r.font.size else None, "bold": r.font.bold}
                             for r in p.runs]} for p in shape.text_frame.paragraphs]
            if shape.has_chart:
                entry["chart_type"] = str(shape.chart.chart_type)
                entry["series"] = [{"name": s.name, "values": list(s.values)} for s in shape.chart.series]
            if shape.has_table:
                entry["table_rows"] = len(shape.table.rows)
                entry["table_columns"] = len(shape.table.columns)
            shapes.append(entry)
        layout_name = ""
        try:
            layout_name = slide.slide_layout.name
        except Exception:
            layout_name = ""
        notes_text = ""
        try:
            if slide.has_notes_slide:
                notes_text = slide.notes_slide.notes_text_frame.text
        except Exception:
            notes_text = ""
        slides.append(
            {
                "index": slide_index,
                "layout": layout_name,
                "shapes": shapes,
                "shape_count": len(shapes),
                "notes": notes_text,
                "layout_index": next((i for i, layout in enumerate(presentation.slide_layouts) if layout.part is slide.slide_layout.part), None),
            }
        )
    payload: Dict[str, Any] = {
        "slides": slides if max_slides <= 0 else slides[:max_slides],
        "slide_count": len(slides),
        "slide_size_emu": {
            "width": presentation.slide_width,
            "height": presentation.slide_height,
        },
    }
    if max_slides > 0:
        payload["slides_truncated"] = len(slides) > max_slides
    return payload


def pptx_read_text(path: Path) -> str:
    require_pptx()
    presentation = _open(path)
    blocks: List[str] = []
    for slide_index, slide in enumerate(presentation.slides):
        blocks.append("## Slide %d" % (slide_index + 1))
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                blocks.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    blocks.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(blocks)
