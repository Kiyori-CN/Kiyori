"""XLSX 读取。

读取公式与缓存值时必须分两次 ``load_workbook``：``data_only=True`` 的载入
只用于取缓存值，绝不能在保存路径上使用（见 ops_xlsx）。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional

from ..protocol import OfficeError

try:
    import openpyxl  # type: ignore
    from openpyxl.utils import get_column_letter  # type: ignore

    OPENPYXL_AVAILABLE = True
except Exception:  # pragma: no cover
    OPENPYXL_AVAILABLE = False


def require_openpyxl() -> None:
    if not OPENPYXL_AVAILABLE:
        raise OfficeError(
            "E_ENV_MISSING",
            "XLSX 处理需要 openpyxl（Tier1）",
            detail="openpyxl is not importable",
            remedy="调用 office_env_setup 安装 Tier1 组件，或先安装 openpyxl",
        )


def _load(path: Path, *, data_only: bool = False, keep_vba: bool = False):
    try:
        return openpyxl.load_workbook(
            str(path), data_only=data_only, keep_vba=keep_vba, read_only=False
        )
    except Exception as exc:
        raise OfficeError(
            "E_DOC_CORRUPT",
            "XLSX 打开失败",
            detail="%s: %s" % (type(exc).__name__, exc),
            remedy="确认文件是未加密的 .xlsx/.xlsm；.xls 旧格式需先 office_convert",
        ) from exc


def is_xlsm(path: Path) -> bool:
    return path.suffix.lower() == ".xlsm"


def xlsx_info(path: Path) -> Dict[str, Any]:
    require_openpyxl()
    workbook = _load(path, keep_vba=is_xlsm(path))
    sheets: List[Dict[str, Any]] = []
    for sheet in workbook.worksheets:
        sheets.append(
            {
                "name": sheet.title,
                "index": workbook.worksheets.index(sheet),
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "dimensions": sheet.dimensions,
                "state": sheet.sheet_state,
            }
        )
    named_ranges = []
    try:
        for name, definition in workbook.defined_names.items():
            named_ranges.append({"name": name, "definition": str(definition.value)})
    except Exception:
        named_ranges = []
    external_links = [str(link) for link in getattr(workbook, "_external_links", [])]
    formula_count = 0
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_count += 1
    return {
        "sheets": sheets,
        "sheet_count": len(sheets),
        "named_ranges": named_ranges,
        "external_links": external_links,
        "has_macros": is_xlsm(path),
        "formula_cells": formula_count,
    }


def _normalize_range(sheet, cell_range: Optional[str]) -> str:
    if not cell_range:
        return "A1:%s%d" % (get_column_letter(max(1, sheet.max_column)), max(1, sheet.max_row))
    return cell_range


def xlsx_read(
    path: Path,
    *,
    sheet_name: Optional[str] = None,
    cell_range: Optional[str] = None,
    max_rows: int = 500,
    with_formulas: bool = True,
) -> Dict[str, Any]:
    """按范围读，同时返回公式与缓存值。"""

    require_openpyxl()
    workbook = _load(path, keep_vba=is_xlsm(path))
    cached = _load(path, data_only=True, keep_vba=is_xlsm(path)) if with_formulas else None

    names = [sheet.title for sheet in workbook.worksheets]
    if sheet_name:
        if sheet_name not in names:
            raise OfficeError(
                "E_INPUT_SCHEMA",
                "工作表不存在: %s" % sheet_name,
                detail="available sheets: %s" % ", ".join(names),
            )
        targets = [workbook[sheet_name]]
    else:
        targets = list(workbook.worksheets)

    result_sheets: List[Dict[str, Any]] = []
    for sheet in targets:
        cell_range = _normalize_range(sheet, cell_range)
        rows: List[List[Dict[str, Any]]] = []
        for row_index, row in enumerate(sheet[cell_range], start=1):
            if row_index > max_rows:
                break
            rendered: List[Dict[str, Any]] = []
            for cell in row:
                value = cell.value
                entry: Dict[str, Any] = {"ref": cell.coordinate, "value": value}
                if isinstance(value, str) and value.startswith("="):
                    entry["formula"] = value
                    if cached is not None:
                        cached_sheet = cached[sheet.title]
                        entry["cached_value"] = cached_sheet[cell.coordinate].value
                rendered.append(entry)
            rows.append(rendered)
        result_sheets.append(
            {
                "name": sheet.title,
                "range": cell_range,
                "rows": rows,
                "row_count": len(rows),
                "truncated": sheet.max_row > max_rows,
            }
        )
    return {"sheets": result_sheets, "sheet_names": names}
