"""Excel 原生可编辑图表；直接引用单元格，保存后仍可随数据更新。"""

from typing import Any, Dict

from ..paths import artifact, atomic_save
from ..protocol import OfficeError, register
from ..readers.xlsx_reader import require_openpyxl
from .write import _open_workbook, _prepare, _resolve_sheet, cache_warnings


@register("xlsx_chart", schema="xlsx_chart", engine="openpyxl",
          next_actions=["xlsx_recalc", "office_validate", "office_render_preview"])
def xlsx_chart(args: Dict[str, Any]) -> Dict[str, Any]:
    require_openpyxl()
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference, ScatterChart, Series
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, range_boundaries

    _, directory, source, output = _prepare(args)
    if source is None:
        raise OfficeError("E_INPUT_SCHEMA", "xlsx_chart 需要 path")
    workbook = _open_workbook(source)
    sheet = _resolve_sheet(workbook, args.get("sheet_name"))
    try:
        bounds = range_boundaries(args["data_range"])
        min_col, min_row, max_col, max_row = bounds
        if None in bounds or not (1 <= min_col < max_col <= 16384 and 1 <= min_row < max_row <= 1048576):
            raise ValueError("至少需要表头、一个数据行和两列")
        anchor = args.get("anchor", "E2")
        column, row = coordinate_from_string(anchor)
        if row < 1 or row > 1048576 or column_index_from_string(column) > 16384:
            raise ValueError("anchor 越界")
    except (ValueError, TypeError) as exc:
        raise OfficeError("E_INPUT_SCHEMA", "data_range/anchor 无效：范围须含表头，首列为分类或 X 值") from exc
    if (max_row - min_row + 1) * (max_col - min_col + 1) > 100000:
        raise OfficeError("E_BUDGET_EXCEEDED", "图表数据超过 100000 单元格，请先聚合或缩小范围")
    for col in range(min_col + 1, max_col + 1):
        if not isinstance(sheet.cell(min_row, col).value, str) or not sheet.cell(min_row, col).value.strip():
            raise OfficeError("E_INPUT_SCHEMA", "每个数据系列需要非空文本表头")
    kind = args.get("chart_type", "column")
    if kind == "pie" and max_col - min_col != 1:
        raise OfficeError("E_INPUT_SCHEMA", "饼图只能有一列分类和一列数值，不能静默丢弃其他系列")
    factories = {"column": BarChart, "bar": BarChart, "line": LineChart, "pie": PieChart, "scatter": ScatterChart}
    if kind not in factories:
        raise OfficeError("E_INPUT_SCHEMA", "chart_type 必须是 column/bar/line/pie/scatter")
    chart = factories[kind]()
    if kind in ("column", "bar"):
        chart.type = "col" if kind == "column" else "bar"
    categories = Reference(sheet, min_col=min_col, min_row=min_row + 1, max_row=max_row)
    if kind == "scatter":
        # X 值须可计算；文本分类轴应使用 line，不能让散点图静默忽略整条系列。
        for row_index in range(min_row + 1, max_row + 1):
            value = sheet.cell(row_index, min_col).value
            if isinstance(value, bool) or not (isinstance(value, (int, float)) or isinstance(value, str) and value.startswith("=")):
                raise OfficeError("E_INPUT_SCHEMA", "散点图首列必须是数字或公式")
        for col in range(min_col + 1, max_col + 1):
            values = Reference(sheet, min_col=col, min_row=min_row, max_row=max_row)
            chart.series.append(Series(values, categories, title_from_data=True))
    else:
        chart.add_data(Reference(sheet, min_col=min_col + 1, max_col=max_col, min_row=min_row, max_row=max_row), titles_from_data=True)
        chart.set_categories(categories)
    chart.title = args.get("title") or None
    chart.style = args.get("style", 10)
    chart.width = args.get("width_cm", 18)
    chart.height = args.get("height_cm", 10)
    sheet.add_chart(chart, anchor)
    atomic_save(workbook, output)
    workbook.close()
    return {"artifacts": [artifact(output, role="in_place" if args.get("in_place") else "output")],
            "warnings": cache_warnings(workbook),
            "data": {"chart_type": kind, "sheet": sheet.title, "data_range": args["data_range"],
                     "series_count": len(chart.series), "anchor": anchor, "work_dir": str(directory)}}
