"""公式策略：openpyxl 写公式的三条硬规则。

1. 溢出数组函数（XLOOKUP/FILTER/SORT/UNIQUE/SEQUENCE）禁止写入：
   openpyxl 没有 spill 元数据，LibreOffice 重算后只有左上角有值，而错误计数仍是 0，
   属于「静默错误」，比报错更危险。
2. TEXTJOIN/CONCAT/IFS/SWITCH/MAXIFS/MINIFS 必须写成 ``_xlfn.`` 前缀。
3. ``data_only=True`` 载入后保存会永久丢失全部公式，保存路径必须拦截。
"""

from __future__ import annotations

import re
from typing import Any, Dict, List, Optional

from ..protocol import OfficeError

SPILL_FUNCTIONS = (
    "XLOOKUP",
    "FILTER",
    "SORT",
    "SORTBY",
    "UNIQUE",
    "SEQUENCE",
    "RANDARRAY",
    "TEXTSPLIT",
    "TOCOL",
    "TOROW",
    "WRAPCOLS",
    "WRAPROWS",
    "EXPAND",
    "TAKE",
    "DROP",
    "CHOOSEROWS",
    "CHOOSECOLS",
    "HSTACK",
    "VSTACK",
    "GROUPBY",
    "PIVOTBY",
)

NEEDS_XLFN_PREFIX = (
    "TEXTJOIN",
    "CONCAT",
    "IFS",
    "SWITCH",
    "MAXIFS",
    "MINIFS",
)

_FUNCTION_CALL = re.compile(r"(?<![A-Za-z0-9_.])([A-Za-z][A-Za-z0-9_.]*)\s*\(")


def normalize_formula(formula: str) -> str:
    """校验并归一化公式：拒绝溢出函数，补齐 _xlfn. 前缀。"""

    if not isinstance(formula, str) or not formula.startswith("="):
        raise OfficeError(
            "E_INPUT_SCHEMA",
            "公式必须以 = 开头",
            detail=repr(formula)[:200],
        )
    for match in _FUNCTION_CALL.finditer(formula):
        raw_name = match.group(1)
        name = raw_name.upper()
        bare = name[6:] if name.startswith("_XLFN.") else name
        if bare in SPILL_FUNCTIONS:
            raise OfficeError(
                "E_INPUT_SCHEMA",
                "禁止写入溢出数组函数 %s" % bare,
                detail=formula[:200],
                remedy=(
                    "溢出函数在 openpyxl 产物中只有左上角有值且错误计数仍为 0，"
                    "属于静默错误；改用 xlsx_aggregate 预聚合，或用 SUMIFS/INDEX/MATCH 等价写法"
                ),
            )
        if bare in NEEDS_XLFN_PREFIX and not raw_name.upper().startswith("_XLFN."):
            formula = formula.replace(
                match.group(0), match.group(0).replace(raw_name, "_xlfn." + raw_name, 1), 1
            )
    return formula


def audit_formulas(path) -> Dict[str, Any]:
    """扫描工作簿里的公式，返回溢出函数与缺少缓存值的单元格清单。"""

    import openpyxl  # type: ignore

    workbook = openpyxl.load_workbook(str(path), data_only=False, keep_vba=str(path).lower().endswith(".xlsm"))
    cached = openpyxl.load_workbook(str(path), data_only=True, keep_vba=str(path).lower().endswith(".xlsm"))
    formula_cells: List[Dict[str, Any]] = []
    spill_cells: List[Dict[str, Any]] = []
    missing_cache: List[Dict[str, Any]] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if not isinstance(value, str) or not value.startswith("="):
                    continue
                upper = value.upper()
                formula_cells.append({"sheet": sheet.title, "cell": cell.coordinate})
                for name in SPILL_FUNCTIONS:
                    if name + "(" in upper:
                        spill_cells.append(
                            {
                                "sheet": sheet.title,
                                "cell": cell.coordinate,
                                "function": name,
                                "formula": value,
                            }
                        )
                        break
                cached_value = cached[sheet.title][cell.coordinate].value
                if cached_value is None:
                    missing_cache.append(
                        {"sheet": sheet.title, "cell": cell.coordinate, "formula": value}
                    )
    return {
        "formula_count": len(formula_cells),
        "spill_functions": spill_cells,
        "cells_without_cached_value": missing_cache,
        "missing_cache_count": len(missing_cache),
    }
