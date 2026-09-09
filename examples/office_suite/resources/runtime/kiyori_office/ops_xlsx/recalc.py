"""xlsx_recalc：LibreOffice 重算并回写。

openpyxl 写出的公式没有缓存值；未经重算，``data_only=True``、pandas 与大多数
预览器读到的都是 ``None``。因此任何写入公式的产物交付前都必须经过本命令，
且 ``total_errors`` 必须为 0。
"""

from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any, Dict, List

from ..env import require_libreoffice
from ..paths import atomic_output, artifact, resolve_output_path, resolve_path, resolve_task_id, task_dir
from ..protocol import OfficeError, engine_version, register
from ..readers.xlsx_reader import require_openpyxl
from .formula import spill_functions

RECALC_PROFILE_XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
  <item oor:path="/org.openoffice.Office.Calc/Formula/Load">
    <prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop>
    <prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop>
  </item>
</oor:items>
"""

def _prepare_profile(directory: Path) -> Path:
    profile = directory / "tmp" / "lo-profile"
    user_dir = profile / "user"
    user_dir.mkdir(parents=True, exist_ok=True)
    (user_dir / "registrymodifications.xcu").write_text(RECALC_PROFILE_XCU, encoding="utf-8")
    return profile


def _run_soffice(soffice: str, profile: Path, source: Path, target_dir: Path) -> Dict[str, Any]:
    command = [
        soffice,
        "--headless",
        "--norestore",
        "--nolockcheck",
        "--nodefault",
        "-env:UserInstallation=%s" % profile.resolve().as_uri(),
        "--convert-to",
        "xlsx:Calc MS Excel 2007 XML",
        "--outdir",
        str(target_dir),
        str(source),
    ]
    env = dict(os.environ)
    env.setdefault("HOME", str(profile.parent))
    completed = subprocess.run(
        command,
        capture_output=True,
        text=True,
        timeout=600,
        env=env,
        check=False,
    )
    return {
        "command": " ".join(command),
        "exit_code": completed.returncode,
        "stdout_tail": (completed.stdout or "")[-2000:],
        "stderr_tail": (completed.stderr or "")[-2000:],
    }


def _audit(path: Path) -> Dict[str, Any]:
    import openpyxl  # type: ignore

    workbook = openpyxl.load_workbook(str(path), data_only=False)
    cached = openpyxl.load_workbook(str(path), data_only=True)
    total_formulas = 0
    error_cells: List[Dict[str, Any]] = []
    missing_cache: List[Dict[str, Any]] = []
    spill_cells: List[Dict[str, Any]] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if not isinstance(value, str) or not value.startswith("="):
                    continue
                total_formulas += 1
                cached_value = cached[sheet.title][cell.coordinate].value
                if cached_value is None:
                    missing_cache.append(
                        {"sheet": sheet.title, "cell": cell.coordinate, "formula": value}
                    )
                elif cached[sheet.title][cell.coordinate].data_type == "e":
                    error_cells.append(
                        {
                            "sheet": sheet.title,
                            "cell": cell.coordinate,
                            "formula": value,
                            "error": cached_value,
                        }
                    )
                for name in spill_functions(value):
                    if name:
                        spill_cells.append(
                            {
                                "sheet": sheet.title,
                                "cell": cell.coordinate,
                                "function": name,
                            }
                        )
                        break
    workbook.close()
    cached.close()
    return {
        "total_formulas": total_formulas,
        "total_errors": len(error_cells),
        "error_cells": error_cells,
        "cells_without_cached_value": missing_cache,
        "missing_cache_count": len(missing_cache),
        "spill_function_cells": spill_cells,
    }


@register(
    "xlsx_recalc",
    schema="xlsx_recalc",
    engine="libreoffice",
    next_actions=["office_validate", "office_render_preview"],
)
def xlsx_recalc(args: Dict[str, Any]) -> Dict[str, Any]:
    require_openpyxl()
    # 重算走 Calc 组件：只装 Writer 的机器上 soffice 存在但重算必然失败。
    soffice = require_libreoffice(purpose="xlsx 公式重算", component="calc")["path"]
    task_id = resolve_task_id(args)
    directory = task_dir(task_id)
    source = resolve_path(args.get("path"), args=args, field="path", must_exist=True)
    if source.suffix.lower() != ".xlsx":
        raise OfficeError("E_FORMAT_UNSUPPORTED", "重算只支持 .xlsx；不把含宏工作簿转换成无宏文件")
    output = resolve_output_path(
        args.get("output_path") or (str(source) if args.get("in_place") else None),
        args=args,
        field="output_path",
        default_dir=directory / "out",
        default_name=source.name,
        overwrite=bool(args.get("overwrite")),
        in_place=bool(args.get("in_place")),
    )

    if output.suffix.lower() != ".xlsx":
        raise OfficeError("E_INPUT_SCHEMA", "重算 output_path 必须是 .xlsx")
    # 每次调用独立目录，避免 soffice 返回 0 却未产物时误读上次结果。
    target_dir = Path(tempfile.mkdtemp(prefix="recalc-", dir=directory / "tmp"))
    profile = _prepare_profile(target_dir)
    try:
        run = _run_soffice(soffice, profile, source, target_dir)
    except subprocess.TimeoutExpired as exc:
        raise OfficeError(
            "E_TIMEOUT",
            "LibreOffice 重算超时",
            detail=str(exc),
            remedy="缩小工作簿范围后重试，或提高 timeoutMs",
        ) from exc
    if run["exit_code"] != 0:
        raise OfficeError(
            "E_ENGINE_FAILED",
            "LibreOffice 重算失败",
            detail="exit=%s stderr=%s" % (run["exit_code"], run["stderr_tail"]),
            remedy="检查 LibreOffice 安装与 profile 目录权限；不要自动切换引擎",
        )
    recalculated = target_dir / source.name
    if not recalculated.is_file():
        raise OfficeError(
            "E_ENGINE_FAILED",
            "LibreOffice 未产出重算结果",
            detail="expected=%s stdout=%s" % (recalculated, run["stdout_tail"]),
        )

    audit = _audit(recalculated)
    if audit["spill_function_cells"]:
        raise OfficeError(
            "E_VALIDATION_FAILED",
            "产物包含禁止的溢出数组函数，重算结果不可信",
            detail="cells=%s" % audit["spill_function_cells"][:10],
            data=audit,
            remedy="改用 SUMIFS/INDEX/MATCH 等 Excel 2007 级函数后重新生成",
        )

    if audit["total_errors"] or audit["missing_cache_count"]:
        raise OfficeError("E_VALIDATION_FAILED", "公式重算存在错误或缺失缓存，未发布产物", data=audit,
                          remedy="按 error_cells 和 cells_without_cached_value 修正工作簿后重算")
    with atomic_output(output) as temporary:
        shutil.copy2(recalculated, temporary)
    audit["bytes"] = output.stat().st_size
    audit["work_dir"] = str(directory)
    return {
        "artifacts": [artifact(output, role="in_place" if args.get("in_place") else "output")],
        "data": audit,
        "engine_version": engine_version("libreoffice"),
        "next_actions": (
            ["office_validate", "office_render_preview"]
            if audit["total_errors"] == 0 and audit["missing_cache_count"] == 0
            else ["xlsx_write"]
        ),
    }
