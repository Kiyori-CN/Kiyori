"""xlsx_write / xlsx_format / xlsx_sheet / xlsx_table / xlsx_export。"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Dict, List

from ..paths import atomic_save, artifact, resolve_output_path, resolve_path, resolve_task_id, task_dir
from ..protocol import OfficeError, engine_version, register
from ..readers.xlsx_reader import is_xlsm, require_openpyxl
from .formula import normalize_formula

def _open_workbook(source: Path):
    """打开可写工作簿。

    保存路径上禁止 ``data_only=True``：openpyxl 以缓存值载入后保存会永久丢掉
    全部公式。这里显式提供 ``data_only`` 形参，让调用方看见并触发拦截。
    """

    return open_workbook_for_write(source)


def open_workbook_for_write(source: Path, *, data_only: bool = False):
    if data_only:
        raise OfficeError(
            "E_INPUT_SCHEMA",
            "禁止用 data_only=True 载入后保存工作簿",
            detail="openpyxl data_only 载入只保留缓存值，保存会永久丢失全部公式",
            remedy="读取缓存值请用 xlsx_read；写入请用默认（data_only=False）载入",
        )
    import openpyxl  # type: ignore

    require_openpyxl()
    try:
        return openpyxl.load_workbook(str(source), keep_vba=is_xlsm(source))
    except Exception as exc:
        raise OfficeError(
            "E_DOC_CORRUPT",
            "XLSX 打开失败",
            detail="%s: %s" % (type(exc).__name__, exc),
        ) from exc


def _prepare(args: Dict[str, Any]):
    task_id = resolve_task_id(args)
    directory = task_dir(task_id)
    raw_path = args.get("path")
    source = resolve_path(raw_path, args=args, field="path", must_exist=True) if raw_path else None
    default_name = source.name if source else str(args.get("file_name") or "workbook.xlsx")
    in_place = bool(args.get("in_place"))
    output = resolve_output_path(
        args.get("output_path") or (str(source) if in_place and source else None),
        args=args,
        field="output_path",
        default_dir=directory / "out",
        default_name=default_name,
        overwrite=bool(args.get("overwrite")),
        in_place=in_place,
    )
    if output.suffix.lower() not in (".xlsx", ".xlsm"):
        raise OfficeError(
            "E_INPUT_SCHEMA",
            "output_path 必须是 .xlsx 或 .xlsm",
            detail=str(output),
        )
    if source and source.suffix.lower() != output.suffix.lower():
        raise OfficeError("E_INPUT_SCHEMA", "编辑必须保留工作簿扩展名，不能隐式移除宏或伪装格式")
    if source is None and output.suffix.lower() == ".xlsm":
        raise OfficeError("E_INPUT_SCHEMA", "新建无宏工作簿请使用 .xlsx；.xlsm 需要原始宏模板")
    return task_id, directory, source, output


def _resolve_sheet(workbook, name):
    """在已有工作簿中定位工作表；找不到即报错（新建语义见 xlsx_write）。"""

    if not name:
        return workbook.active
    if name in workbook.sheetnames:
        return workbook[name]
    raise OfficeError(
        "E_INPUT_SCHEMA",
        "工作表不存在: %s" % name,
        detail="available=%s" % ", ".join(workbook.sheetnames),
    )


def cache_warnings(workbook):
    # 任意 openpyxl 保存都会失去旧缓存，连仅调整列宽也不例外。按已有单元格
    # 检查，不展开稀疏表的巨大空白矩形，并把下一步明确交给 Agent。
    if any(cell.data_type == "f" for sheet in workbook.worksheets for cell in sheet._cells.values()):
        return [{"code": "FORMULA_CACHE_INVALIDATED", "message": "本次保存已清除公式缓存；完成全部写入、格式、图表及工作表操作后，再执行 xlsx_recalc 并严格校验。"}]
    return []


SHEET_TITLE_ILLEGAL = set('[]:*?/\\')


def _validate_sheet_title(name: str) -> str:
    """按 Excel 约束校验工作表名，避免 openpyxl 抛出难以诊断的异常。"""

    title = name.strip()
    if not title:
        raise OfficeError("E_INPUT_SCHEMA", "sheet_name 不能为空")
    if len(title) > 31:
        raise OfficeError(
            "E_INPUT_SCHEMA",
            "sheet_name 超过 31 个字符",
            detail="length=%d" % len(title),
        )
    illegal = sorted(SHEET_TITLE_ILLEGAL.intersection(title))
    if illegal:
        raise OfficeError(
            "E_INPUT_SCHEMA",
            "sheet_name 含非法字符",
            detail="illegal=%s" % "".join(illegal),
            remedy="工作表名不能包含 [ ] : * ? / \\",
        )
    return title


@register(
    "xlsx_write",
    schema="xlsx_write",
    engine="openpyxl",
    next_actions=["xlsx_recalc", "xlsx_format", "office_validate"],
)
def xlsx_write(args: Dict[str, Any]) -> Dict[str, Any]:
    require_openpyxl()
    task_id, directory, source, output = _prepare(args)
    if source is not None:
        workbook = _open_workbook(source)
        sheet = _resolve_sheet(workbook, args.get("sheet_name"))
    else:
        # 新建工作簿：sheet_name 表示「默认表叫什么」，直接重命名默认表，
        # 不能走 _resolve_sheet 的「查找已存在表」语义（历史故障：
        # 新建时传 sheet_name 报「工作表不存在」）。
        workbook = __import__("openpyxl").Workbook()
        sheet = workbook.active
        requested_name = args.get("sheet_name")
        if requested_name:
            sheet.title = _validate_sheet_title(str(requested_name))

    cells = args.get("cells")
    rows = args.get("rows")
    if cells is None and rows is None:
        raise OfficeError(
            "E_INPUT_SCHEMA",
            "xlsx_write 需要 cells 或 rows",
            remedy="cells=[{cell, value|formula}]，rows=[[...]]（可用 start_cell 指定起点）",
        )
    # rows 是基础矩阵；更精确的 cells 最后覆盖（含显式 null 清空）。
    # 在内存中汇总最终单元格后再写入，返回计数描述实际产物而非赋值次数。
    pending = {}
    assignments = 0
    from openpyxl.utils.cell import coordinate_to_tuple

    def coordinate(ref):
        if not isinstance(ref, str) or not re.fullmatch(r"[A-Za-z]{1,3}[1-9][0-9]*", ref):
            raise OfficeError("E_INPUT_SCHEMA", "单元格地址必须是 A1 形式：%r" % ref)
        row, col = coordinate_to_tuple(ref.upper())
        if row > 1048576 or col > 16384:
            raise OfficeError("E_INPUT_SCHEMA", "单元格地址超出 Excel 范围：%s" % ref)
        return row, col

    def collect(key, value):
        nonlocal assignments
        if isinstance(value, str) and value.startswith("="):
            value = normalize_formula(value)
        pending[key] = value
        assignments += 1
        if assignments > 100000:
            raise OfficeError("E_BUDGET_EXCEEDED", "单次写入超过 100000 个单元格，请分批写入")

    if isinstance(rows, list):
        start_row, start_column = coordinate(args.get("start_cell") or "A1")
        for row_offset, row in enumerate(rows):
            if not isinstance(row, list):
                raise OfficeError("E_INPUT_SCHEMA", "rows[%d] 必须是数组" % row_offset)
            if start_row + row_offset > 1048576 or start_column + len(row) - 1 > 16384:
                raise OfficeError("E_INPUT_SCHEMA", "rows 超出 Excel 行列范围")
            for column_offset, value in enumerate(row):
                collect((start_row + row_offset, start_column + column_offset), value)
    if isinstance(cells, list):
        for index, item in enumerate(cells):
            if not isinstance(item, dict) or ("value" in item) == ("formula" in item):
                raise OfficeError("E_INPUT_SCHEMA", "cells[%d] 必须且只能提供 value 或 formula" % index)
            key = coordinate(item.get("cell"))
            if "formula" in item:
                if not isinstance(item["formula"], str):
                    raise OfficeError("E_INPUT_SCHEMA", "cells[%d].formula 必须是字符串" % index)
                collect(key, normalize_formula(item["formula"]))
            else:
                collect(key, item["value"])
    if not pending:
        raise OfficeError("E_INPUT_SCHEMA", "cells/rows 没有可写入单元格")
    for (row, col), value in pending.items():
        sheet.cell(row, col).value = value
    written = len(pending)
    formula_count = sum(isinstance(v, str) and v.startswith("=") for v in pending.values())

    output.parent.mkdir(parents=True, exist_ok=True)
    atomic_save(workbook, output)
    return {
        "artifacts": [artifact(output, role="in_place" if args.get("in_place") else "output")],
        "data": {
            "cells_written": written,
            "overwritten_in_batch": assignments - written,
            "write_order": "rows_then_cells",
            "formula_cells": formula_count,
            "work_dir": str(directory),
            "sheet": sheet.title,
        },
        "warnings": cache_warnings(workbook),
        "engine_version": engine_version("openpyxl"),
        "next_actions": (
            ["xlsx_recalc", "office_validate"]
            if formula_count or source is not None
            else ["xlsx_format", "office_validate"]
        ),
    }


def _column_index(ref: str) -> int:
    from openpyxl.utils import column_index_from_string  # type: ignore

    match = re.match(r"([A-Z]+)", ref.upper())
    if not match:
        raise OfficeError("E_INPUT_SCHEMA", "start_cell 非法: %s" % ref)
    return column_index_from_string(match.group(1))


@register(
    "xlsx_format",
    schema="xlsx_format",
    engine="openpyxl",
    next_actions=["xlsx_recalc", "office_validate"],
)
def xlsx_format(args: Dict[str, Any]) -> Dict[str, Any]:
    require_openpyxl()
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # type: ignore
    from openpyxl.utils import get_column_letter  # type: ignore

    task_id, directory, source, output = _prepare(args)
    if source is None:
        raise OfficeError("E_INPUT_SCHEMA", "xlsx_format 需要 path")
    workbook = _open_workbook(source)
    sheet = _resolve_sheet(workbook, args.get("sheet_name"))
    applied: List[str] = []
    from openpyxl.utils.cell import range_boundaries
    selected_range = args.get("range") or sheet.dimensions
    try:
        min_col, min_row, max_col, max_row = range_boundaries(selected_range)
        if None in (min_col, min_row, max_col, max_row) or not (1 <= min_col <= max_col <= 16384 and 1 <= min_row <= max_row <= 1048576):
            raise ValueError("需要有限矩形范围")
    except (ValueError, TypeError) as exc:
        raise OfficeError("E_INPUT_SCHEMA", "range 必须是 A1:D20 这样的有限矩形范围") from exc
    if (max_col - min_col + 1) * (max_row - min_row + 1) > 100000:
        raise OfficeError("E_BUDGET_EXCEEDED", "格式范围超过 100000 个单元格，请缩小 range")

    def selected_rows():
        return sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)

    number_format = args.get("number_format")
    if number_format:
        for row in selected_rows():
            for cell in row:
                cell.number_format = str(number_format)
        applied.append("number_format")

    font = args.get("font")
    if isinstance(font, dict):
        from copy import copy
        for row in selected_rows():
            for cell in row:
                # 仅更新明确指定的属性，改字号不能抹掉原有字体、颜色或粗体。
                style = copy(cell.font)
                for key in ("name", "size", "bold", "italic", "color"):
                    if key in font:
                        setattr(style, key, font[key])
                cell.font = style
        applied.append("font")

    fill = args.get("fill")
    if isinstance(fill, dict) and fill.get("color"):
        pattern = PatternFill(fill_type=fill.get("type") or "solid", fgColor=str(fill["color"]))
        for row in selected_rows():
            for cell in row:
                cell.fill = pattern
        applied.append("fill")

    widths = args.get("column_widths")
    if isinstance(widths, dict):
        for key, value in widths.items():
            sheet.column_dimensions[str(key).upper()].width = float(value)
        applied.append("column_widths")

    heights = args.get("row_heights")
    if isinstance(heights, dict):
        for key, value in heights.items():
            sheet.row_dimensions[int(key)].height = float(value)
        applied.append("row_heights")

    freeze = args.get("freeze_panes")
    if freeze:
        sheet.freeze_panes = str(freeze)
        applied.append("freeze_panes")

    autofilter = args.get("auto_filter")
    if autofilter:
        sheet.auto_filter.ref = str(autofilter)
        applied.append("auto_filter")

    output.parent.mkdir(parents=True, exist_ok=True)
    atomic_save(workbook, output)
    return {
        "artifacts": [artifact(output, role="in_place" if args.get("in_place") else "output")],
        "data": {"applied": applied, "work_dir": str(directory)},
        "warnings": cache_warnings(workbook),
        "engine_version": engine_version("openpyxl"),
    }


@register(
    "xlsx_sheet",
    schema="xlsx_sheet",
    engine="openpyxl",
    next_actions=["xlsx_write", "xlsx_recalc", "office_validate"],
)
def xlsx_sheet(args: Dict[str, Any]) -> Dict[str, Any]:
    require_openpyxl()
    task_id, directory, source, output = _prepare(args)
    if source is None:
        raise OfficeError("E_INPUT_SCHEMA", "xlsx_sheet 需要 path")
    workbook = _open_workbook(source)
    operation = str(args.get("operation") or "")
    name = str(args.get("name") or "")
    if operation not in ("create", "delete", "rename", "move", "copy"):
        raise OfficeError(
            "E_INPUT_SCHEMA",
            "operation 必须是 create/delete/rename/move/copy",
        )
    if operation == "create":
        if not name:
            raise OfficeError("E_INPUT_SCHEMA", "create 需要 name")
        name = _validate_sheet_title(name)
        if name.casefold() in {title.casefold() for title in workbook.sheetnames}:
            raise OfficeError("E_INPUT_SCHEMA", "工作表已存在: %s" % name)
        index = args.get("index")
        if index is not None and (isinstance(index, bool) or not isinstance(index, int) or not 0 <= index <= len(workbook.sheetnames)):
            raise OfficeError("E_INPUT_SCHEMA", "create index 超出工作表范围")
        workbook.create_sheet(name, index=index)
    else:
        sheet = _resolve_sheet(workbook, name)
        if operation == "delete":
            if len(workbook.sheetnames) == 1:
                raise OfficeError("E_INPUT_SCHEMA", "不能删除唯一的工作表")
            workbook.remove(sheet)
        elif operation == "rename":
            new_name = str(args.get("new_name") or "")
            if not new_name:
                raise OfficeError("E_INPUT_SCHEMA", "rename 需要 new_name")
            new_name = _validate_sheet_title(new_name)
            if any(s is not sheet and s.title.casefold() == new_name.casefold() for s in workbook.worksheets):
                raise OfficeError("E_INPUT_SCHEMA", "工作表已存在: %s" % new_name)
            sheet.title = new_name
        elif operation == "move":
            index = args.get("index")
            if isinstance(index, bool) or not isinstance(index, int) or not 0 <= index < len(workbook.sheetnames):
                raise OfficeError("E_INPUT_SCHEMA", "move 需要整数 index")
            current = workbook.sheetnames.index(sheet.title)
            workbook.move_sheet(sheet, offset=index - current)
        else:
            target_name = str(args.get("new_name") or ("%s-copy" % sheet.title))
            target_name = _validate_sheet_title(target_name)
            if target_name.casefold() in {title.casefold() for title in workbook.sheetnames}:
                raise OfficeError("E_INPUT_SCHEMA", "目标工作表已存在: %s" % target_name)
            if sheet._charts or sheet._images or sheet.tables:
                raise OfficeError("E_FORMAT_UNSUPPORTED", "工作表复制暂不支持包含图表、图片或 Excel Table 的源表",
                                  remedy="先复制基础数据表，再在目标表显式创建图表和表格；避免交付静默丢失对象的副本")
            workbook.copy_worksheet(sheet).title = target_name
    output.parent.mkdir(parents=True, exist_ok=True)
    atomic_save(workbook, output)
    return {
        "artifacts": [artifact(output, role="in_place" if args.get("in_place") else "output")],
        "data": {
            "operation": operation,
            "sheets": list(workbook.sheetnames),
            "work_dir": str(directory),
        },
        "warnings": cache_warnings(workbook),
        "engine_version": engine_version("openpyxl"),
    }


@register(
    "xlsx_table",
    schema="xlsx_table",
    engine="openpyxl",
    next_actions=["xlsx_recalc", "office_validate"],
)
def xlsx_table(args: Dict[str, Any]) -> Dict[str, Any]:
    require_openpyxl()
    from openpyxl.utils import get_column_letter  # type: ignore
    from openpyxl.worksheet.table import Table, TableStyleInfo  # type: ignore

    task_id, directory, source, output = _prepare(args)
    if source is None:
        raise OfficeError("E_INPUT_SCHEMA", "xlsx_table 需要 path")
    workbook = _open_workbook(source)
    sheet = _resolve_sheet(workbook, args.get("sheet_name"))
    ref = str(args.get("range") or "")
    if not ref:
        ref = "A1:%s%d" % (get_column_letter(sheet.max_column), sheet.max_row)
    name = str(args.get("table_name") or "Table1")
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name=str(args.get("style") or "TableStyleMedium2"),
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    output.parent.mkdir(parents=True, exist_ok=True)
    atomic_save(workbook, output)
    return {
        "artifacts": [artifact(output, role="in_place" if args.get("in_place") else "output")],
        "data": {"table": name, "range": ref, "work_dir": str(directory)},
        "warnings": cache_warnings(workbook),
        "engine_version": engine_version("openpyxl"),
    }
