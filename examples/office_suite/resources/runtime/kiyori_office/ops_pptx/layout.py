"""显式、可编辑的幻灯片布局；不靠静默缩字掩盖内容溢出。"""

import math
import re

from ..protocol import OfficeError


def measure(value, field, minimum=0, maximum=200):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or not minimum <= value <= maximum:
        raise OfficeError("E_INPUT_SCHEMA", "%s 必须是 %s 到 %s 的有限数字" % (field, minimum, maximum))
    return value


def rgb(value):
    from pptx.dml.color import RGBColor
    if not isinstance(value, str) or not re.fullmatch(r"[a-fA-F0-9]{6}", value):
        raise OfficeError("E_INPUT_SCHEMA", "颜色必须为六位 RGB 十六进制字符串")
    return RGBColor.from_string(value)


def set_font(font, config):
    from pptx.util import Pt
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml.ns import qn
    if "font_name" in config:
        name = config["font_name"]
        if not isinstance(name, str) or not name.strip():
            raise OfficeError("E_INPUT_SCHEMA", "font_name 必须是非空字符串")
        font.name = name
        props = font._rPr
        east = props.find(qn("a:ea"))
        if east is None:
            east = OxmlElement("a:ea")
            props.append(east)
        east.set("typeface", name)
    if "size_pt" in config:
        font.size = Pt(measure(config["size_pt"], "size_pt", 1, 400))
    for key in ("bold", "italic"):
        if key in config:
            if not isinstance(config[key], bool):
                raise OfficeError("E_INPUT_SCHEMA", "%s 必须是 boolean" % key)
            setattr(font, key, config[key])
    if "color_rgb" in config:
        font.color.rgb = rgb(config["color_rgb"])


def text_frame(frame, paragraphs, config):
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.util import Pt, Cm
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml.ns import qn
    if not isinstance(paragraphs, list) or not paragraphs:
        raise OfficeError("E_INPUT_SCHEMA", "paragraphs 必须是非空数组")
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    anchors = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
    anchor = config.get("vertical_alignment", "top")
    if anchor not in anchors:
        raise OfficeError("E_INPUT_SCHEMA", "vertical_alignment 必须是 top/middle/bottom")
    frame.vertical_anchor = anchors[anchor]
    if "margin_cm" in config:
        margin = Cm(measure(config["margin_cm"], "margin_cm", 0, 5))
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = margin
    aligns = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
    for index, item in enumerate(paragraphs):
        if isinstance(item, str):
            item = {"text": item}
        if not isinstance(item, dict) or not isinstance(item.get("text"), str):
            raise OfficeError("E_INPUT_SCHEMA", "paragraphs[] 必须是字符串或含 text 的对象")
        values = {**config, **item}
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item["text"]
        alignment = values.get("alignment", "left")
        if alignment not in aligns:
            raise OfficeError("E_INPUT_SCHEMA", "alignment 必须是 left/center/right/justify")
        paragraph.alignment = aligns[alignment]
        level = item.get("level", 0)
        if isinstance(level, bool) or not isinstance(level, int) or not 0 <= level <= 8:
            raise OfficeError("E_INPUT_SCHEMA", "level 必须为 0 到 8 的整数")
        paragraph.level = level
        for key in ("space_before_pt", "space_after_pt"):
            if key in values:
                setattr(paragraph, key[:-3], Pt(measure(values[key], key, 0, 400)))
        if "line_spacing" in values:
            paragraph.line_spacing = measure(values["line_spacing"], "line_spacing", 0.5, 10)
        if "bullet" in item:
            if not isinstance(item["bullet"], bool):
                raise OfficeError("E_INPUT_SCHEMA", "bullet 必须是 boolean")
            props = paragraph._p.get_or_add_pPr()
            for child in list(props):
                if child.tag in {qn("a:buNone"), qn("a:buChar"), qn("a:buAutoNum")}:
                    props.remove(child)
            bullet = OxmlElement("a:buChar" if item["bullet"] else "a:buNone")
            if item["bullet"]:
                bullet.set("char", "•")
                # 留出悬挂缩进，换行不压住项目符号。
                props.set("marL", str(int(Cm(0.5 * (level + 1)))))
                props.set("indent", str(-int(Cm(0.3))))
            props.append(bullet)
        set_font(paragraph.font, values)
        for run in paragraph.runs:
            set_font(run.font, values)


def add_elements(prs, slide, elements):
    from pptx.util import Cm
    from pptx.enum.shapes import MSO_SHAPE
    if not isinstance(elements, list):
        raise OfficeError("E_INPUT_SCHEMA", "elements 必须是数组")
    for index, spec in enumerate(elements):
        if not isinstance(spec, dict):
            raise OfficeError("E_INPUT_SCHEMA", "elements[] 必须是对象")
        kind = spec.get("type", "text")
        common = {"type", "name", "left_cm", "top_cm", "width_cm", "height_cm"}
        typography = {"font_name", "size_pt", "bold", "italic", "color_rgb", "alignment", "vertical_alignment",
                      "margin_cm", "space_before_pt", "space_after_pt", "line_spacing"}
        allowed = {
            "text": common | typography | {"text", "paragraphs", "fill_rgb"},
            "shape": common | typography | {"text", "paragraphs", "fill_rgb", "line_rgb", "shape"},
            "table": common | typography | {"rows"},
            "chart": common | typography | {"chart_type", "categories", "series", "title"},
        }
        if kind not in allowed or set(spec) - allowed[kind]:
            raise OfficeError("E_INPUT_SCHEMA", "elements[%d] 类型或字段不支持：%s" % (index, sorted(set(spec) - allowed.get(kind, set()))))
        if "text" in spec and "paragraphs" in spec:
            raise OfficeError("E_INPUT_SCHEMA", "text 与 paragraphs 只能提供一个")
        box = [Cm(measure(spec.get(key), key, 0.01 if key in ("width_cm", "height_cm") else 0))
               for key in ("left_cm", "top_cm", "width_cm", "height_cm")]
        left, top, width, height = box
        if left + width > prs.slide_width or top + height > prs.slide_height:
            raise OfficeError("E_INPUT_SCHEMA", "elements[%d] 超出幻灯片边界" % index)
        if kind == "text":
            shape = slide.shapes.add_textbox(*box)
            text_frame(shape.text_frame, spec.get("paragraphs", [spec.get("text", "")]), spec)
        elif kind == "shape":
            kinds = {"rectangle": MSO_SHAPE.RECTANGLE, "rounded_rectangle": MSO_SHAPE.ROUNDED_RECTANGLE, "oval": MSO_SHAPE.OVAL}
            name = spec.get("shape", "rectangle")
            if name not in kinds:
                raise OfficeError("E_INPUT_SCHEMA", "shape 必须是 rectangle/rounded_rectangle/oval")
            shape = slide.shapes.add_shape(kinds[name], *box)
            if "text" in spec or "paragraphs" in spec:
                text_frame(shape.text_frame, spec.get("paragraphs", [spec.get("text", "")]), spec)
        elif kind == "table":
            rows = spec.get("rows")
            if not isinstance(rows, list) or not rows or not all(isinstance(row, list) and row for row in rows) or len({len(row) for row in rows}) != 1:
                raise OfficeError("E_INPUT_SCHEMA", "table.rows 必须是非空矩形二维数组")
            if len(rows) * len(rows[0]) > 2000:
                raise OfficeError("E_BUDGET_EXCEEDED", "PPT 单表超过 2000 个单元格")
            shape = slide.shapes.add_table(len(rows), len(rows[0]), *box)
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    text_frame(shape.table.cell(r, c).text_frame, ["" if value is None else str(value)], spec)
        elif kind == "chart":
            from pptx.chart.data import CategoryChartData
            from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
            types = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
                     "line": XL_CHART_TYPE.LINE, "pie": XL_CHART_TYPE.PIE}
            chart_type = spec.get("chart_type", "column")
            categories, series = spec.get("categories"), spec.get("series")
            if chart_type not in types or not isinstance(categories, list) or not categories or not all(isinstance(c, str) for c in categories):
                raise OfficeError("E_INPUT_SCHEMA", "chart 需要 column/bar/line/pie 类型和非空文本 categories")
            if not isinstance(series, list) or not series or chart_type == "pie" and len(series) != 1:
                raise OfficeError("E_INPUT_SCHEMA", "chart 需要非空 series；饼图只能有一个系列")
            if len(categories) * len(series) > 10000:
                raise OfficeError("E_BUDGET_EXCEEDED", "PPT 图表超过 10000 个数据点，请先聚合")
            data = CategoryChartData()
            data.categories = categories
            names = set()
            for item in series:
                if not isinstance(item, dict) or not isinstance(item.get("name"), str) or not item["name"].strip() or item["name"] in names:
                    raise OfficeError("E_INPUT_SCHEMA", "series[].name 必须是唯一的非空文本")
                names.add(item["name"])
                values = item.get("values")
                if not isinstance(values, list) or len(values) != len(categories):
                    raise OfficeError("E_INPUT_SCHEMA", "series[].values 长度必须等于 categories")
                for value in values:
                    if value is not None:
                        measure(value, "series[].values", -1e100, 1e100)
                data.add_series(item["name"], values)
            shape = slide.shapes.add_chart(types[chart_type], *box, data)
            chart = shape.chart
            chart.has_legend = len(series) > 1 or chart_type == "pie"
            if chart.has_legend:
                chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                chart.legend.include_in_layout = False
                set_font(chart.legend.font, spec)
            if "title" in spec:
                chart.has_title = True
                text_frame(chart.chart_title.text_frame, [spec["title"]], spec)
            if chart_type != "pie":
                set_font(chart.category_axis.tick_labels.font, spec)
                set_font(chart.value_axis.tick_labels.font, spec)
        else:
            raise OfficeError("E_INPUT_SCHEMA", "elements[].type 必须是 text/shape/table/chart")
        if "name" in spec:
            if not isinstance(spec["name"], str) or not spec["name"].strip() or any(s != shape and s.name == spec["name"] for s in slide.shapes):
                raise OfficeError("E_INPUT_SCHEMA", "元素 name 必须为页内唯一的非空字符串")
            shape.name = spec["name"]
        if "fill_rgb" in spec and kind not in ("table", "chart"):
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(spec["fill_rgb"])
        if kind == "shape" and "line_rgb" in spec:
            shape.line.color.rgb = rgb(spec["line_rgb"])
        elif kind == "shape":
            shape.line.fill.background()
