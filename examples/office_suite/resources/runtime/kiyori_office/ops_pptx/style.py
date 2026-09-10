"""创建与增量编辑共用的原生 DrawingML 样式。"""

from ..protocol import OfficeError
from .layout import measure, rgb, set_font


def checked(value, allowed, field):
    if not isinstance(value, dict) or set(value) - set(allowed):
        raise OfficeError("E_INPUT_SCHEMA", "%s 含不支持的字段或不是对象" % field)
    return value


def boolean(value, field):
    if not isinstance(value, bool):
        raise OfficeError("E_INPUT_SCHEMA", "%s 必须是 boolean" % field)
    return value


def solid(fill, color):
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def shape_style(shape, config):
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml.ns import qn
    from pptx.util import Pt
    checked(config, {"fill_rgb", "line_rgb", "line_width_pt", "opacity", "gradient", "shadow", "corner_radius"}, "style")
    if "fill_rgb" in config and "gradient" in config:
        raise OfficeError("E_INPUT_SCHEMA", "fill_rgb 与 gradient 不能同时指定")
    if "fill_rgb" in config:
        solid(shape.fill, config["fill_rgb"])
    if "gradient" in config:
        gradient = checked(config["gradient"], {"colors", "angle"}, "gradient")
        colors = gradient.get("colors")
        if not isinstance(colors, list) or len(colors) != 2:
            raise OfficeError("E_INPUT_SCHEMA", "gradient.colors 需要两个 RGB 色值")
        shape.fill.gradient()
        for stop, color in zip(shape.fill.gradient_stops, colors):
            stop.color.rgb = rgb(color)
        shape.fill.gradient_angle = measure(gradient.get("angle", 0), "gradient.angle", 0, 360)
    if "opacity" in config:
        opacity = measure(config["opacity"], "opacity", 0, 1)
        colors = shape._element.xpath("./p:spPr/a:solidFill/* | ./p:spPr/a:gradFill/a:gsLst/a:gs/*")
        if not colors:
            raise OfficeError("E_INPUT_SCHEMA", "opacity 需要显式实色或渐变填充")
        for color in colors:
            for child in list(color):
                if child.tag == qn("a:alpha"):
                    color.remove(child)
            alpha = OxmlElement("a:alpha")
            alpha.set("val", str(round(opacity * 100000)))
            color.append(alpha)
    if "line_rgb" in config:
        shape.line.color.rgb = rgb(config["line_rgb"])
    if "line_width_pt" in config:
        shape.line.width = Pt(measure(config["line_width_pt"], "line_width_pt", 0, 30))
    if "corner_radius" in config:
        from pptx.enum.shapes import MSO_SHAPE
        if shape.auto_shape_type != MSO_SHAPE.ROUNDED_RECTANGLE:
            raise OfficeError("E_INPUT_SCHEMA", "corner_radius 仅用于 rounded_rectangle")
        shape.adjustments[0] = measure(config["corner_radius"], "corner_radius", 0, 0.5)
    if "shadow" in config:
        values = checked(config["shadow"], {"color_rgb", "opacity", "blur_pt", "distance_pt", "angle"}, "shadow")
        props = shape._element.spPr
        for node in list(props):
            if node.tag in {qn("a:effectLst"), qn("a:effectDag")}:
                props.remove(node)
        effects = OxmlElement("a:effectLst")
        shadow = OxmlElement("a:outerShdw")
        shadow.set("blurRad", str(int(Pt(measure(values.get("blur_pt", 4), "blur_pt", 0, 100)))))
        shadow.set("dist", str(int(Pt(measure(values.get("distance_pt", 2), "distance_pt", 0, 100)))))
        shadow.set("dir", str(round(measure(values.get("angle", 45), "angle", 0, 360) * 60000)))
        shadow.set("rotWithShape", "0")
        color = OxmlElement("a:srgbClr")
        color.set("val", str(rgb(values.get("color_rgb", "000000"))))
        alpha = OxmlElement("a:alpha")
        alpha.set("val", str(round(measure(values.get("opacity", 0.2), "shadow.opacity", 0, 1) * 100000)))
        color.append(alpha)
        shadow.append(color)
        effects.append(shadow)
        props.append(effects)


def table_style(table, config):
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml.ns import qn
    from pptx.util import Cm, Pt
    checked(config, {"header_fill_rgb", "header_color_rgb", "body_fill_rgb", "band_fill_rgb", "border_rgb",
                     "border_width_pt", "font_name", "size_pt", "color_rgb", "column_widths_cm", "row_heights_cm"}, "table_style")
    # 显式单元格颜色覆盖 Office 默认蓝主题，同时保留真实表格语义。
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            key = "header_fill_rgb" if r == 0 else "band_fill_rgb" if r % 2 == 0 and "band_fill_rgb" in config else "body_fill_rgb"
            if key in config:
                solid(cell.fill, config[key])
            font = {k: v for k, v in config.items() if k in {"font_name", "size_pt", "color_rgb"}}
            if r == 0 and "header_color_rgb" in config:
                font.update(color_rgb=config["header_color_rgb"], bold=True)
            for paragraph in cell.text_frame.paragraphs:
                set_font(paragraph.font, font)
                for run in paragraph.runs:
                    set_font(run.font, font)
            if "border_rgb" in config:
                props = cell._tc.get_or_add_tcPr()
                for side in ("lnL", "lnR", "lnT", "lnB"):
                    for old in list(props):
                        if old.tag == qn("a:" + side):
                            props.remove(old)
                    line = OxmlElement("a:" + side)
                    line.set("w", str(int(Pt(measure(config.get("border_width_pt", 0.5), "border_width_pt", 0, 20)))))
                    fill = OxmlElement("a:solidFill")
                    color = OxmlElement("a:srgbClr")
                    color.set("val", str(rgb(config["border_rgb"])))
                    fill.append(color)
                    line.append(fill)
                    props.insert_element_before(line, "a:lnTlToBr", "a:lnBlToTr", "a:cell3D", "a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill", "a:headers", "a:extLst")
    for key, entries in (("column_widths_cm", table.columns), ("row_heights_cm", table.rows)):
        if key in config:
            values = config[key]
            if not isinstance(values, list) or len(values) != len(entries):
                raise OfficeError("E_INPUT_SCHEMA", "%s 数量必须与表格一致" % key)
            for entry, value in zip(entries, values):
                setattr(entry, "width" if key.startswith("column") else "height", Cm(measure(value, key, 0.01, 100)))


def chart_style(chart, config):
    from pptx.enum.chart import XL_LEGEND_POSITION, XL_DATA_LABEL_POSITION
    checked(config, {"series_colors", "point_colors", "legend_position", "legend_font_size_pt", "data_labels",
                     "gridlines", "number_format", "minimum_scale", "maximum_scale", "font_name", "size_pt", "color_rgb",
                     "background_rgb", "border_rgb", "plot_fill_rgb"}, "chart_style")
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.dml.fill import FillFormat
    for key, parent in (("background_rgb", chart._chartSpace), ("plot_fill_rgb", chart._chartSpace.chart.plotArea)):
        if key in config:
            props = parent.find(qn("c:spPr"))
            if props is None:
                props = OxmlElement("c:spPr")
                parent.insert_element_before(props, "c:txPr", "c:externalData", "c:printSettings", "c:userShapes", "c:extLst")
            fill = FillFormat.from_fill_parent(props)
            fill.background() if config[key] == "none" else solid(fill, config[key])
    if "border_rgb" in config:
        props = chart._chartSpace.find(qn("c:spPr"))
        if props is None:
            props = OxmlElement("c:spPr")
            chart._chartSpace.insert_element_before(props, "c:txPr", "c:externalData", "c:printSettings", "c:userShapes", "c:extLst")
        fill = FillFormat.from_fill_parent(props.get_or_add_ln())
        fill.background() if config["border_rgb"] == "none" else solid(fill, config["border_rgb"])
    if "series_colors" in config:
        colors = config["series_colors"]
        if not isinstance(colors, list) or len(colors) != len(chart.series):
            raise OfficeError("E_INPUT_SCHEMA", "series_colors 数量必须与系列一致")
        for series, color in zip(chart.series, colors):
            solid(series.format.fill, color)
            series.format.line.color.rgb = rgb(color)
    if "point_colors" in config:
        colors = config["point_colors"]
        if len(chart.series) != 1 or not isinstance(colors, list) or len(colors) != len(chart.series[0].points):
            raise OfficeError("E_INPUT_SCHEMA", "point_colors 需要单系列且与数据点数量一致")
        for point, color in zip(chart.series[0].points, colors):
            solid(point.format.fill, color)
    if "legend_position" in config:
        positions = {"bottom": XL_LEGEND_POSITION.BOTTOM, "top": XL_LEGEND_POSITION.TOP,
                     "left": XL_LEGEND_POSITION.LEFT, "right": XL_LEGEND_POSITION.RIGHT}
        value = config["legend_position"]
        if value not in {*positions, "none"}:
            raise OfficeError("E_INPUT_SCHEMA", "legend_position 必须为 top/bottom/left/right/none")
        chart.has_legend = value != "none"
        if chart.has_legend:
            chart.legend.position = positions[value]
            chart.legend.include_in_layout = False
    if chart.has_legend:
        font = {**config}
        if "legend_font_size_pt" in config:
            font["size_pt"] = config["legend_font_size_pt"]
        set_font(chart.legend.font, font)
    if "data_labels" in config:
        labels = checked(config["data_labels"], {"show_value", "show_percentage", "show_category", "position", "number_format", "size_pt", "color_rgb"}, "data_labels")
        positions = {"center": XL_DATA_LABEL_POSITION.CENTER, "inside_end": XL_DATA_LABEL_POSITION.INSIDE_END,
                     "outside_end": XL_DATA_LABEL_POSITION.OUTSIDE_END, "best_fit": XL_DATA_LABEL_POSITION.BEST_FIT}
        for plot in chart.plots:
            plot.has_data_labels = True
            target = plot.data_labels
            for key in ("show_value", "show_percentage", "show_category"):
                if key in labels:
                    setattr(target, "show_category_name" if key == "show_category" else key, boolean(labels[key], key))
            if "position" in labels:
                if labels["position"] not in positions:
                    raise OfficeError("E_INPUT_SCHEMA", "data_labels.position 不支持")
                target.position = positions[labels["position"]]
            if "number_format" in labels:
                target.number_format = labels["number_format"]
                target.number_format_is_linked = False
            set_font(target.font, {**config, **labels})
    has_axes = bool(chart._chartSpace.xpath(".//c:valAx"))
    axis_keys = {"gridlines", "number_format", "minimum_scale", "maximum_scale"}
    if not has_axes and set(config) & axis_keys:
        raise OfficeError("E_INPUT_SCHEMA", "无坐标轴图表不接受坐标轴样式")
    if has_axes:
        axis = chart.value_axis
        if "gridlines" in config:
            axis.has_major_gridlines = boolean(config["gridlines"], "gridlines")
        if "number_format" in config:
            axis.tick_labels.number_format = config["number_format"]
            axis.tick_labels.number_format_is_linked = False
        for key in ("minimum_scale", "maximum_scale"):
            if key in config:
                setattr(axis, key, measure(config[key], key, -1e100, 1e100))
        if axis.minimum_scale is not None and axis.maximum_scale is not None and axis.minimum_scale >= axis.maximum_scale:
            raise OfficeError("E_INPUT_SCHEMA", "minimum_scale 必须小于 maximum_scale")
        set_font(axis.tick_labels.font, config)
        set_font(chart.category_axis.tick_labels.font, config)


def replace_chart_data(chart, config):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    checked(config, {"categories", "series"}, "chart_data")
    external = chart._chartSpace.xpath('./c:externalData')
    if external:
        from pptx.oxml.ns import qn
        relation = chart.part.rels[external[0].get(qn('r:id'))]
        if relation.is_external or relation.target_part.content_type != 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
            raise OfficeError('E_FORMAT_UNSUPPORTED', '不改写外部链接或非标准 XLSX 图表数据源')
    if len(chart.plots) != 1 or chart.chart_type not in {XL_CHART_TYPE.COLUMN_CLUSTERED, XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.LINE, XL_CHART_TYPE.PIE}:
        raise OfficeError("E_FORMAT_UNSUPPORTED", "数据编辑仅支持单一柱状/条形/折线/饼图，保留复杂图表原状")
    categories, series = config.get("categories"), config.get("series")
    if not isinstance(categories,list) or not categories or not all(isinstance(value,str) for value in categories):
        raise OfficeError("E_INPUT_SCHEMA", "chart_data.categories 必须为非空字符串数组")
    if not isinstance(series,list) or not series or len(series)*len(categories)>10000:
        raise OfficeError("E_INPUT_SCHEMA", "chart_data.series 必须非空且最多10000数据点")
    if chart.chart_type == XL_CHART_TYPE.PIE and len(series)!=1:
        raise OfficeError("E_INPUT_SCHEMA", "饼图只能有一个系列")
    data=CategoryChartData()
    data.categories=categories
    names=set()
    for item in series:
        checked(item,{"name","values"},"chart_data.series[]")
        name=item.get("name")
        values=item.get("values")
        if not isinstance(name,str) or not name.strip() or name in names or not isinstance(values,list) or len(values)!=len(categories):
            raise OfficeError("E_INPUT_SCHEMA", "系列名称须唯一非空，数值数量须等于类别数量")
        names.add(name)
        for value in values:
            if value is not None:
                measure(value,"chart_data.values",-1e100,1e100)
        data.add_series(name,values)
    chart.replace_data(data)
