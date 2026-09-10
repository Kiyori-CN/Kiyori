"""pptx_edit / pptx_notes / pptx_media / pptx_outline 命令注册。"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List

from ..paths import atomic_save, artifact, resolve_output_path, resolve_path, resolve_task_id, task_dir
from ..protocol import OfficeError, engine_version, register
from ..readers.pptx_reader import pptx_outline, require_pptx
from .slides import _open, _prepare

@register("pptx_outline", schema="pptx_outline", next_actions=["pptx_edit", "pptx_slide"])
def pptx_outline_command(args: Dict[str, Any]) -> Dict[str, Any]:
    source = resolve_path(args.get("path"), args=args, field="path", must_exist=True)
    outline = pptx_outline(source, max_slides=int(args.get("max_slides") or 0))
    if args.get("layout_report"):
        from .quality import slide_report
        prs = _open(source)
        for entry in outline["slides"]:
            entry["layout_report"] = slide_report(prs.slides[entry["index"]], prs)
    return {
        "data": {
            **outline,
            "navigation": [
                {"slide": slide["index"] + 1, "layout": slide["layout"]}
                for slide in outline["slides"]
            ],
        }
    }


def _find_shape(slide, shape_index=None, shape_name=None, shape_id=None):
    shapes = list(slide.shapes)
    if shape_id is not None:
        def walk(items):
            for shape in items:
                yield shape
                if hasattr(shape, "shapes"):
                    yield from walk(shape.shapes)
        matches = [shape for shape in walk(shapes) if shape.shape_id == shape_id]
        if len(matches) != 1:
            raise OfficeError("E_ANCHOR_NOT_FOUND", "shape_id 未唯一命中")
        return matches[0]
    if shape_index is not None:
        if not isinstance(shape_index, int) or isinstance(shape_index, bool):
            raise OfficeError("E_INPUT_SCHEMA", "shape_index 必须是整数")
        if shape_index < 0 or shape_index >= len(shapes):
            raise OfficeError(
                "E_ANCHOR_NOT_FOUND",
                "shape_index 越界: %s" % shape_index,
                detail="shape_count=%d" % len(shapes),
                remedy="先调用 pptx_outline 获取形状列表",
            )
        return shapes[shape_index]
    if shape_name:
        matches = [(index, shape) for index, shape in enumerate(shapes) if shape.name == shape_name]
        if len(matches) > 1:
            raise OfficeError('E_ANCHOR_NOT_FOUND', '形状名命中多处，拒绝模糊编辑',
                              detail='shape_name=%s indices=%s' % (shape_name, [index for index, _ in matches]),
                              remedy='使用 pptx_outline 返回的 shape_index 精确定位。')
        if matches:
            return matches[0][1]
        raise OfficeError(
            "E_ANCHOR_NOT_FOUND",
            "未找到形状: %s" % shape_name,
            remedy="先调用 pptx_outline 获取形状名",
        )
    raise OfficeError("E_INPUT_SCHEMA", "需要 shape_index 或 shape_name")


def _resolve_slide(prs, index: Any):
    if not isinstance(index, int) or isinstance(index, bool):
        raise OfficeError("E_INPUT_SCHEMA", "slide_index 必须是 0-based 整数")
    slides = list(prs.slides)
    if index < 0 or index >= len(slides):
        raise OfficeError(
            "E_ANCHOR_NOT_FOUND",
            "slide_index 越界: %s" % index,
            detail="slide_count=%d" % len(slides),
            remedy="先调用 pptx_outline 获取幻灯片列表",
        )
    return slides[index]


@register(
    "pptx_edit",
    schema="pptx_edit",
    engine="python-pptx",
    next_actions=["office_validate", "office_render_preview"],
)
def pptx_edit(args: Dict[str, Any]) -> Dict[str, Any]:
    require_pptx()
    from pptx.util import Emu, Pt  # type: ignore

    task_id, directory, source, output = _prepare(args)
    prs = _open(source)
    slide = _resolve_slide(prs, args.get("slide_index"))
    operation = str(args.get("operation") or "set_text")
    extended = {"add_elements", "delete", "replace_image", "set_style", "set_table", "set_chart", "z_order", "group", "ungroup", "set_transition", "set_animations"}
    if operation in extended:
        changed = _edit_objects(prs, slide, args)
        output.parent.mkdir(parents=True, exist_ok=True)
        atomic_save(prs, output)
        return {"artifacts": [artifact(output, role="in_place" if args.get("in_place") else "output")],
                "data": {"operation": operation, "changed": changed, "work_dir": str(directory)},
                "engine_version": engine_version("python-pptx")}
    shape = _find_shape(slide, args.get("shape_index"), args.get("shape_name"), args.get("shape_id"))
    if operation not in ("set_text", "append_text", "set_position", "set_size", "set_font"):
        raise OfficeError(
            "E_INPUT_SCHEMA",
            "operation 必须是 set_text/append_text/set_position/set_size/set_font",
        )
    changed: List[str] = []
    if operation in ("set_text", "append_text"):
        if "text" not in args:
            raise OfficeError("E_INPUT_SCHEMA", "%s 需要 text" % operation)
        if not shape.has_text_frame:
            raise OfficeError(
                "E_INPUT_SCHEMA", "目标形状没有文本框", detail=shape.name
            )
        text = str(args.get("text") or "")
        if operation == "set_text":
            # 替换文本保留首段/首 run 的外观；clear() 会移除 run 格式，导致
            # 设计好的字号、字体和颜色突然回退到母版。
            from copy import deepcopy
            first = shape.text_frame.paragraphs[0]
            paragraph_properties = deepcopy(first._p.pPr) if first._p.pPr is not None else None
            run_properties = deepcopy(first.runs[0]._r.rPr) if first.runs and first.runs[0]._r.rPr is not None else None
            shape.text_frame.clear()
            for index, line in enumerate(text.split("\n")):
                paragraph = shape.text_frame.paragraphs[0] if index == 0 else shape.text_frame.add_paragraph()
                paragraph.text = line
                if paragraph_properties is not None:
                    existing = paragraph._p.pPr
                    if existing is not None:
                        paragraph._p.remove(existing)
                    paragraph._p.insert(0, deepcopy(paragraph_properties))
                if paragraph.runs and run_properties is not None:
                    paragraph.runs[0]._r.insert(0, deepcopy(run_properties))
        else:
            paragraph = shape.text_frame.add_paragraph()
            paragraph.text = text
        changed.append(operation)
        if args.get("size_pt"):
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(float(args["size_pt"]))
            changed.append("size_pt")
    elif operation == "set_position":
        if "left_emu" not in args and "top_emu" not in args:
            raise OfficeError("E_INPUT_SCHEMA", "set_position 需要 left_emu/top_emu")
        if args.get("left_emu") is not None:
            shape.left = Emu(int(args["left_emu"]))
            changed.append("left")
        if args.get("top_emu") is not None:
            shape.top = Emu(int(args["top_emu"]))
            changed.append("top")
    elif operation == "set_size":
        if "width_emu" not in args and "height_emu" not in args:
            raise OfficeError("E_INPUT_SCHEMA", "set_size 需要 width_emu/height_emu")
        if args.get("width_emu") is not None:
            shape.width = Emu(int(args["width_emu"]))
            changed.append("width")
        if args.get("height_emu") is not None:
            shape.height = Emu(int(args["height_emu"]))
            changed.append("height")
    else:
        if not shape.has_text_frame:
            raise OfficeError("E_INPUT_SCHEMA", "目标形状没有文本框", detail=shape.name)
        from .layout import set_font
        if not set(args) & {"size_pt", "bold", "italic", "color_rgb", "font_name"}:
            raise OfficeError("E_INPUT_SCHEMA", "set_font 需要字体参数")
        for paragraph in shape.text_frame.paragraphs:
            set_font(paragraph.font, args)
            for run in paragraph.runs:
                set_font(run.font, args)
        changed.append("font")
    output.parent.mkdir(parents=True, exist_ok=True)
    atomic_save(prs, output)
    return {
        "artifacts": [artifact(output, role="in_place" if args.get("in_place") else "output")],
        "data": {"operation": operation, "changed": changed, "work_dir": str(directory)},
        "engine_version": engine_version("python-pptx"),
    }


def _edit_objects(prs, slide, args):
    from .layout import add_elements
    from .style import shape_style, table_style, chart_style
    from .motion import transition, animations
    operation = args["operation"]
    if operation == "add_elements":
        added = add_elements(prs, slide, args.get("elements"), args=args, theme=args.get("theme"))
        return [{"shape_id": shape.shape_id, "name": shape.name} for shape in added]
    if operation == "set_transition":
        transition(slide, args.get("transition"))
        return ["transition"]
    if operation == "set_animations":
        animations(slide, args.get("animations"))
        return ["animations"]
    if operation == "group":
        ids = args.get("shape_ids")
        if not isinstance(ids, list) or len(ids) < 2 or len(set(ids)) != len(ids):
            raise OfficeError("E_INPUT_SCHEMA", "group 需要至少两个不重复的 shape_ids")
        shapes = [s for s in slide.shapes if s.shape_id in ids]
        if len(shapes) != len(ids):
            raise OfficeError("E_ANCHOR_NOT_FOUND", "group 仅接受同一页顶层对象")
        # 非连续层级成组必然改变中间对象的遮挡关系，要求先明确重排图层。
        indices = [list(slide.shapes).index(s) for s in shapes]
        if max(indices) - min(indices) + 1 != len(indices):
            raise OfficeError("E_INPUT_SCHEMA", "分组对象的图层必须连续，请先 z_order")
        tree = slide.shapes._spTree
        position = list(tree).index(shapes[0]._element)
        group = slide.shapes.add_group_shape(shapes)
        tree.remove(group._element)
        tree.insert(position, group._element)
        if args.get("group_name"):
            if any(s.name == args["group_name"] for s in slide.shapes if s.shape_id != group.shape_id):
                raise OfficeError("E_INPUT_SCHEMA", "group_name 已存在")
            group.name = args["group_name"]
        return [{"shape_id": group.shape_id, "name": group.name}]
    shape = _find_shape(slide, args.get("shape_index"), args.get("shape_name"), args.get("shape_id"))
    if operation == "delete":
        ids = {str(shape.shape_id)} | {n.get("id") for n in shape._element.xpath(".//p:cNvPr")}
        if any(node.get("spid") in ids for node in slide._element.xpath(".//p:spTgt")):
            raise OfficeError("E_FORMAT_UNSUPPORTED", "对象被动画引用，请先显式更新 animations")
        if any(node.get("id") in ids for node in slide._element.xpath(".//a:stCxn | .//a:endCxn")):
            raise OfficeError("E_FORMAT_UNSUPPORTED", "对象被连接线引用，删除会破坏连接")
        candidates = _relationship_ids(shape._element)
        shape._element.getparent().remove(shape._element)
        _drop_unused_relationships(slide, candidates)
    elif operation == "replace_image":
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            raise OfficeError("E_INPUT_SCHEMA", "replace_image 仅用于图片")
        image = resolve_path(args.get("image_path"), args=args, field="image_path", must_exist=True)
        _, relationship = slide.part.get_or_add_image_part(str(image))
        old = shape._element.blipFill.blip.rEmbed
        shape._element.blipFill.blip.rEmbed = relationship
        _drop_unused_relationships(slide, {old})
    elif operation == "set_style":
        shape_style(shape, args.get("style"))
    elif operation == "set_table":
        if not shape.has_table:
            raise OfficeError("E_INPUT_SCHEMA", "目标不是表格")
        if "cells" in args:
            cells = args["cells"]
            if not isinstance(cells, list):
                raise OfficeError("E_INPUT_SCHEMA", "cells 必须是数组")
            from .style import checked
            for cell in cells:
                checked(cell, {"row", "column", "text"}, "cells[]")
                r, c = cell.get("row"), cell.get("column")
                if type(r) is not int or type(c) is not int or not 0 <= r < len(shape.table.rows) or not 0 <= c < len(shape.table.columns) or not isinstance(cell.get("text"), str):
                    raise OfficeError("E_INPUT_SCHEMA", "cell 需要有效的 row/column 和 text")
                target = shape.table.cell(r, c)
                if target.is_spanned:
                    raise OfficeError("E_INPUT_SCHEMA", "不能写入被合并遮盖的单元格")
                from .layout import text_frame
                # 只替换选中的单元格，其他单元格/合并信息保持原样。
                from copy import deepcopy
                props = deepcopy(target.text_frame.paragraphs[0]._p)
                target.text = cell["text"]
                old_font = props.find("{http://schemas.openxmlformats.org/drawingml/2006/main}r")
                if old_font is not None and len(old_font) and old_font[0].tag.endswith("}rPr"):
                    for run in target.text_frame.paragraphs[0].runs:
                        run._r.insert(0, deepcopy(old_font[0]))
        if "table_style" in args:
            table_style(shape.table, args["table_style"])
    elif operation == "set_chart":
        if not shape.has_chart:
            raise OfficeError("E_INPUT_SCHEMA", "目标不是图表")
        if "chart_data" not in args and "chart_style" not in args:
            raise OfficeError("E_INPUT_SCHEMA", "set_chart 需要 chart_data/chart_style")
        chart = _isolate_chart(slide, shape)
        if "chart_data" in args:
            from .style import replace_chart_data
            replace_chart_data(chart, args["chart_data"])
        if "chart_style" in args:
            chart_style(chart, args["chart_style"])
    elif operation == "z_order":
        parent = shape._element.getparent()
        siblings = [s for s in parent if s.tag.rsplit("}", 1)[-1] in {"sp", "pic", "graphicFrame", "grpSp", "cxnSp"}]
        order = args.get("z_index")
        if type(order) is not int or not 0 <= order < len(siblings):
            raise OfficeError("E_INPUT_SCHEMA", "z_index 必须是所在组的有效 0-based 图层索引，0 为最底层")
        if siblings[order] is not shape._element:
            old_index = siblings.index(shape._element)
            reference = siblings[order]
            parent.remove(shape._element)
            reference.addprevious(shape._element) if order < old_index else reference.addnext(shape._element)
    elif operation == "ungroup":
        if not hasattr(shape, "shapes"):
            raise OfficeError("E_INPUT_SCHEMA", "目标不是组合对象")
        from pptx.oxml.ns import qn
        xfrm = shape._element.grpSpPr.xfrm
        # 非单位变换的解除组合涉及旋转/翻转/缩放字体，拒绝破坏外观。
        if xfrm.get("rot", "0") != "0" or xfrm.get("flipH", "0") != "0" or xfrm.get("flipV", "0") != "0" or xfrm.ext.cx != xfrm.chExt.cx or xfrm.ext.cy != xfrm.chExt.cy:
            raise OfficeError("E_FORMAT_UNSUPPORTED", "旋转、翻转或缩放后的组合暂不支持无损解除")
        if any(n.get("spid") == str(shape.shape_id) for n in slide._element.xpath(".//p:spTgt")):
            raise OfficeError("E_FORMAT_UNSUPPORTED", "组合被动画引用，请先更新 animations")
        for child in list(shape.shapes):
            child.left += xfrm.off.x - xfrm.chOff.x
            child.top += xfrm.off.y - xfrm.chOff.y
            shape._element.addprevious(child._element)
        shape._element.getparent().remove(shape._element)
    return [operation]


def _relationship_ids(element):
    prefix = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    return {value for node in element.iter() for key, value in node.attrib.items() if key.startswith(prefix)}


def _isolate_chart(slide, shape):
    """既有文档可能共享 chart/workbook；单对象编辑必须先分离可变部件。"""
    from pptx.oxml.ns import qn
    from pptx.opc.packuri import PackURI
    source = shape.chart.part
    package = slide.part.package
    used = {str(part.partname) for part in package.iter_parts()}

    def clone(part, template):
        index = 1
        while template % index in used:
            index += 1
        name = template % index
        used.add(name)
        return type(part).load(PackURI(name), part.content_type, package, part.blob)

    target = clone(source, '/ppt/charts/chart%d.xml')
    mapping = {}
    for old_id, relation in source.rels.items():
        if relation.is_external:
            mapping[old_id] = target.rels.get_or_add_ext_rel(relation.reltype, relation.target_ref)
        else:
            child = relation.target_part
            if relation.reltype.endswith('/package'):
                extension = str(child.partname).rsplit('.', 1)[-1]
                child = clone(child, '/ppt/embeddings/chartData%d.' + extension)
            mapping[old_id] = target.relate_to(child, relation.reltype)
    prefix = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
    for node in target._element.iter():
        for key, value in list(node.attrib.items()):
            if key.startswith(prefix) and value in mapping:
                node.set(key, mapping[value])
    chart_node = shape._element.xpath('.//c:chart')[0]
    old_id = chart_node.get(qn('r:id'))
    chart_node.set(qn('r:id'), slide.part.relate_to(target, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart'))
    _drop_unused_relationships(slide, {old_id})
    return target.chart


def _drop_unused_relationships(slide, candidates):
    # 只删除本次对象曾引用且页内已无引用的关系，保留共享媒体与版式关系。
    remaining = _relationship_ids(slide._element)
    for identity in candidates - remaining:
        if identity in slide.part.rels:
            slide.part.drop_rel(identity)


@register(
    "pptx_notes",
    schema="pptx_notes",
    engine="python-pptx",
    next_actions=["office_validate"],
)
def pptx_notes(args: Dict[str, Any]) -> Dict[str, Any]:
    require_pptx()
    task_id, directory, source, output = _prepare(args)
    prs = _open(source)
    operation = str(args.get("operation") or "read")
    if operation not in ("read", "write"):
        raise OfficeError("E_INPUT_SCHEMA", "operation 必须是 read/write")
    if operation == "read":
        notes = []
        for index, slide in enumerate(prs.slides):
            text = ""
            if slide.has_notes_slide:
                text = slide.notes_slide.notes_text_frame.text
            notes.append({"slide": index, "notes": text})
        return {"data": {"notes": notes, "work_dir": str(directory)}}
    slide = _resolve_slide(prs, args.get("slide_index"))
    slide.notes_slide.notes_text_frame.text = str(args.get("text") or "")
    output.parent.mkdir(parents=True, exist_ok=True)
    atomic_save(prs, output)
    return {
        "artifacts": [artifact(output, role="in_place" if args.get("in_place") else "output")],
        "data": {"slide_index": args.get("slide_index"), "work_dir": str(directory)},
        "engine_version": engine_version("python-pptx"),
    }


@register(
    "pptx_media",
    schema="pptx_media",
    engine="python-pptx",
    next_actions=["office_validate", "office_render_preview"],
)
def pptx_media(args: Dict[str, Any]) -> Dict[str, Any]:
    require_pptx()
    task_id, directory, source, output = _prepare(args)
    prs = _open(source)
    slide = _resolve_slide(prs, args.get("slide_index"))
    image = resolve_path(args.get("image_path"), args=args, field="image_path", must_exist=True)
    left = int(args.get("left_emu") or 0)
    top = int(args.get("top_emu") or 0)
    width = args.get("width_emu")
    height = args.get("height_emu")
    from PIL import Image
    with Image.open(image) as handle:
        ratio = handle.height / handle.width
    if left < 0 or top < 0 or left >= prs.slide_width or top >= prs.slide_height:
        raise OfficeError("E_INPUT_SCHEMA", "图片起点必须位于幻灯片内部")
    if width is None and height is None:
        width = min(int(prs.slide_width * 0.8), prs.slide_width - left,
                    int((prs.slide_height - top) / ratio))
        height = int(width * ratio)
    elif width is None:
        width = int(height / ratio)
    elif height is None:
        height = int(width * ratio)
    if width <= 0 or height <= 0 or left + width > prs.slide_width or top + height > prs.slide_height:
        raise OfficeError("E_INPUT_SCHEMA", "图片显式尺寸超出幻灯片，请调整位置或尺寸")
    slide.shapes.add_picture(str(image), left, top, width=width, height=height)
    output.parent.mkdir(parents=True, exist_ok=True)
    atomic_save(prs, output)
    return {
        "artifacts": [artifact(output, role="in_place" if args.get("in_place") else "output")],
        "data": {"image": str(image), "width_emu": width, "height_emu": height, "work_dir": str(directory)},
        "engine_version": engine_version("python-pptx"),
    }
