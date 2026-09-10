"""按保存后文档与原始部件核验，覆盖新建、既有文件增量编辑及失败原子性。"""

from zipfile import ZipFile

import pytest

from kiyori_office import protocol


def call(command, **args):
    result = protocol.run(command, args)
    assert result['ok'], result
    return result


def element(kind='text', name='title', **kwargs):
    return dict(type=kind, name=name, left_cm=1, top_cm=1, width_cm=8, height_cm=3, **kwargs)


def deck(tmp_path, elements=None):
    path = tmp_path / 'deck.pptx'
    call('pptx_create', output_path=str(path), layout_index=6,
         slides=[{'elements': elements or [element(text='keep me')]}])
    return path


def test_layout_index_and_theme_are_real(tmp_path):
    from pptx import Presentation
    path = tmp_path / 'blank.pptx'
    result = call('pptx_create', output_path=str(path), layout_index=6,
                  theme={'font_name':'Arial','size_pt':24,'color_rgb':'123456'},
                  slides=[{'elements':[element(text='Theme')]}])
    assert result['data']['layouts'] == ['Blank']
    prs = Presentation(path)
    assert prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.size.pt == 24
    result = call('pptx_outline', path=str(path), layout_report=True)
    assert result['data']['slides'][0]['layout_index'] == 6
    assert result['data']['slides'][0]['layout_report']['elements'][0]['shape_id'] == 2


def test_incremental_table_chart_image_and_styles(tmp_path):
    from PIL import Image
    from pptx import Presentation
    image = tmp_path / 'wide.png'
    Image.new('RGB',(400,100),'purple').save(image)
    source = deck(tmp_path)
    original = source.read_bytes()
    output = tmp_path / 'edited.pptx'
    call('pptx_edit', path=str(source), slide_index=0, output_path=str(output), operation='add_elements', elements=[
        element('table','table', rows=[['H','V'],['a',2]], table_style={'header_fill_rgb':'112233','header_color_rgb':'FFFFFF','body_fill_rgb':'EEEEEE','border_rgb':'777777','size_pt':18}),
        element('chart','chart', categories=['a','b'],series=[{'name':'S','values':[1,2]}],chart_style={'series_colors':['AABBCC'],'gridlines':False,'legend_position':'right','data_labels':{'show_value':True},'number_format':'0.0'}),
        element('image','picture', image_path=str(image), fit='cover'),
        element('shape','panel',shape='rounded_rectangle',style={'gradient':{'colors':['123456','ABCDEF'],'angle':45},'opacity':0.6,'shadow':{'blur_pt':4},'corner_radius':0.2})])
    assert source.read_bytes() == original
    prs = Presentation(output)
    shapes = prs.slides[0].shapes
    assert len(shapes) == 5 and shapes[0].text == 'keep me'
    assert str(shapes[1].table.cell(0,0).fill.fore_color.rgb) == '112233'
    assert str(shapes[2].chart.series[0].format.fill.fore_color.rgb) == 'AABBCC'
    assert not shapes[2].chart.value_axis.has_major_gridlines
    assert shapes[3].crop_left > 0
    call('office_validate', path=str(output), strict=True)
    call('pptx_edit',path=str(output),slide_index=0,shape_name='table',operation='set_table',
         cells=[{'row':1,'column':1,'text':'42'}],in_place=True)
    assert Presentation(output).slides[0].shapes[1].table.cell(1,1).text == '42'


def test_group_z_order_stable_id_and_ungroup(tmp_path):
    from pptx import Presentation
    source = deck(tmp_path,[element(name='a',text='A'),element(name='b',text='B'),element(name='c',text='C')])
    call('pptx_edit',path=str(source),slide_index=0,operation='group',shape_ids=[2,3],group_name='card',in_place=True)
    call('pptx_edit',path=str(source),slide_index=0,operation='set_text',shape_id=2,text='changed',in_place=True)
    prs = Presentation(source)
    assert prs.slides[0].shapes[0].shapes[0].text == 'changed'
    report = call('pptx_outline',path=str(source),layout_report=True)['data']['slides'][0]['layout_report']
    assert any(e['shape_id']==2 and e['parent_id'] is not None for e in report['elements'])
    call('pptx_edit',path=str(source),slide_index=0,operation='ungroup',shape_name='card',in_place=True)
    call('pptx_edit',path=str(source),slide_index=0,operation='z_order',shape_id=2,z_index=2,in_place=True)
    assert [s.shape_id for s in Presentation(source).slides[0].shapes] == [3,4,2]


def test_motion_survives_edit_and_copy(tmp_path):
    from pptx import Presentation
    path = deck(tmp_path)
    call('pptx_edit',path=str(path),slide_index=0,operation='set_transition',transition={'effect':'push','direction':'l','advance_after_ms':3000},in_place=True)
    call('pptx_edit',path=str(path),slide_index=0,operation='set_animations',animations=[
        {'shape_name':'title','effect':'fade','trigger':'on_click'},
        {'shape_name':'title','effect':'wipe','trigger':'after_previous','exit':True}],in_place=True)
    with ZipFile(path) as archive:
        xml = archive.read('ppt/slides/slide1.xml')
    call('pptx_slide',path=str(path),operation='duplicate',index=0,in_place=True)
    prs = Presentation(path)
    for slide in prs.slides:
        assert slide._element.xpath('./p:transition/p:push')[0].get('dir') == 'l'
        assert len(slide._element.xpath('.//p:animEffect')) == 2
        assert slide._element.xpath('.//p:cond[@delay="500"]')
        assert len(slide._element.xpath('.//p:cTn[@nodeType="mainSeq"]/p:childTnLst/p:par')) == 1
        assert len(slide._element.xpath('.//p:attrName[text()="style.visibility"]')) == 2
    before = path.read_bytes()
    result = protocol.run('pptx_edit',dict(path=str(path),slide_index=0,shape_id=2,operation='delete',in_place=True))
    assert not result['ok'] and path.read_bytes() == before
    call('office_validate',path=str(path),strict=True)


@pytest.mark.parametrize('operation,extra', [
    ('set_style',{'style':{'opacity':2}}), ('set_style',{'style':{'gradient':{'colors':['BAD']}}}),
    ('set_chart',{'chart_style':{}}), ('group',{'shape_ids':[2,999]}),
    ('set_animations',{'animations':[{'shape_name':'missing'}]}),
    ('set_transition',{'transition':{'effect':'morph'}})])
def test_bad_edits_preserve_input(tmp_path,operation,extra):
    path = deck(tmp_path)
    before = path.read_bytes()
    result = protocol.run('pptx_edit',dict(path=str(path),slide_index=0,shape_id=2,operation=operation,in_place=True,**extra))
    assert not result['ok'] and path.read_bytes() == before


OMML = '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:f><m:num><m:r><m:t>a</m:t></m:r></m:num><m:den><m:r><m:t>b</m:t></m:r></m:den></m:f></m:oMath>'


def test_insert_paper_blocks_preserves_existing_math_and_image(tmp_path):
    from docx import Document
    from PIL import Image
    image = tmp_path/'plot.png'
    Image.new('RGB',(100,100),'green').save(image)
    path = tmp_path/'paper.docx'
    call('docx_create',output_path=str(path),spec={'blocks':[{'type':'paragraph','text':'anchor'}, {'type':'formula','omml':OMML}]})
    call('docx_edit',path=str(path),operation='insert_blocks_after',anchor={'index':0},in_place=True,blocks=[
        {'type':'paragraph','text':'new'}, {'type':'table','rows':[['a','b'],['c','d']]},
        {'type':'formula','omml':OMML,'number':1},
        {'type':'image','image_path':str(image),'width_cm':3,'caption':'Figure 1','alt_text':'plot'}])
    doc = Document(path)
    assert doc.paragraphs[0].text == 'anchor' and doc.paragraphs[1].text == 'new'
    assert len(doc.element.xpath('.//m:oMath')) == 2
    assert len(doc.inline_shapes) == 1 and len(doc.tables) == 1
    assert doc.inline_shapes[0]._inline.docPr.get('descr') == 'plot'
    call('office_validate',path=str(path),strict=True)


def test_ppt_native_formula(tmp_path):
    path=deck(tmp_path,[element('formula','equation',omml=OMML)])
    with ZipFile(path) as archive:
        xml=archive.read('ppt/slides/slide1.xml')
    assert b'oMath' in xml and b'2010/main' in xml
    call('office_validate',path=str(path),strict=True)


def test_word_native_chart_has_independent_workbook_and_three_lines(tmp_path):
    from docx import Document
    path=tmp_path/'charts.docx'
    call('docx_create',output_path=str(path),spec={'blocks':[
        {'type':'chart','categories':['a','b'],'series':[{'name':'Series','values':[3,7]}],
         'chart_style':{'series_colors':['7755CC'],'gridlines':False},'caption':'Chart 1'},
        {'type':'chart','categories':['a','b'],'series':[{'name':'Other','values':[4,8]}]},
        {'type':'table','rows':[['a','b'],[1,2]],'border_style':'three_line'}]})
    with ZipFile(path) as archive:
        assert len([n for n in archive.namelist() if n.startswith('word/embeddings/')])==2
        assert b'7755CC' in archive.read('word/charts/chart1.xml')
    doc=Document(path)
    assert len(doc.inline_shapes)==2
    assert doc.tables[0]._tbl.xpath('./w:tblPr/w:tblBorders/w:insideV')[0].get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val')=='nil'
    call('office_validate',path=str(path),strict=True)


def test_dry_run_and_overflow_are_honest(tmp_path):
    data=call('pptx_measure_text',text='Dense text ' * 100,width_cm=3,height_cm=1,size_pt=24)['data']
    assert data['estimated_overflow'] and data['visual_verification_required']
    path=deck(tmp_path,[element(text='Dense text '*100,size_pt=24)])
    report=call('pptx_outline',path=str(path),layout_report=True)['data']['slides'][0]['layout_report']
    assert any(i['code']=='POSSIBLE_TEXT_OVERFLOW' for i in report['issues'])


def test_word_edit_native_objects_isolated_from_shared_chart(tmp_path):
    from docx import Document
    from copy import deepcopy
    from PIL import Image
    from io import BytesIO
    from openpyxl import load_workbook
    image=tmp_path/'old.png'
    replacement=tmp_path/'new.png'
    Image.new('RGB',(100,100),'red').save(image)
    Image.new('RGB',(100,100),'green').save(replacement)
    path=tmp_path/'editable.docx'
    call('docx_create',output_path=str(path),spec={'blocks':[
        {'type':'chart','categories':['A'],'series':[{'name':'S','values':[1]}],'chart_style':{'series_colors':['7755CC']}},
        {'type':'formula','omml':OMML}, {'type':'image','image_path':str(image),'width_cm':2}]})
    doc=Document(path)
    doc.paragraphs[-1]._p.addnext(deepcopy(doc.paragraphs[0]._p))
    doc.save(path)
    call('docx_edit',path=str(path),anchor={'index':0},operation='set_chart',in_place=True,
         chart_data={'categories':['A','B'],'series':[{'name':'S','values':[4,9]}]})
    call('docx_edit',path=str(path),anchor={'index':1},operation='set_formula',omml=OMML.replace('>a<','>x<'),in_place=True)
    call('docx_edit',path=str(path),anchor={'index':2},operation='replace_image',image_path=str(replacement),in_place=True)
    with ZipFile(path) as archive:
        charts=[n for n in archive.namelist() if n.startswith('word/charts/') and n.endswith('.xml')]
        assert len(charts)==2
        assert all(b'7755CC' in archive.read(n) for n in charts)
        workbooks=[n for n in archive.namelist() if n.startswith('word/embeddings/')]
        values=[load_workbook(BytesIO(archive.read(n)),data_only=True).active['B2'].value for n in workbooks]
        assert sorted(values)==[1,4]
        media=[n for n in archive.namelist() if n.startswith('word/media/')]
        assert len(media)==1 and archive.read(media[0])==replacement.read_bytes()
    assert '>x<' in Document(path).element.xml
    call('office_validate',path=str(path),strict=True)


def test_ppt_chart_data_and_deleted_media_are_not_stale(tmp_path):
    from pptx import Presentation
    from PIL import Image
    image=tmp_path/'picture.png'
    Image.new('RGB',(100,100),'red').save(image)
    path=deck(tmp_path,[element('chart','chart',categories=['A'],series=[{'name':'S','values':[1]}]),element('image','picture',image_path=str(image))])
    call('pptx_edit',path=str(path),slide_index=0,shape_name='chart',operation='set_chart',
         chart_data={'categories':['B'],'series':[{'name':'S','values':[8]}]},in_place=True)
    call('pptx_edit',path=str(path),slide_index=0,shape_name='picture',operation='delete',in_place=True)
    assert list(Presentation(path).slides[0].shapes[0].chart.series[0].values)==[8]
    with ZipFile(path) as archive:
        assert not any(n.startswith('ppt/media/') for n in archive.namelist())
    call('office_validate',path=str(path),strict=True)


def test_all_icons_are_editable_paths_without_theme_shadow(tmp_path):
    from kiyori_office.ops_pptx.icons import PATHS
    from pptx import Presentation
    path=deck(tmp_path,[element('icon',name,icon=name) for name in PATHS])
    for group in Presentation(path).slides[0].shapes:
        assert len(group.shapes)>0
        assert all(not child.shadow.inherit for child in group.shapes)


def test_pdf_images_and_wrapping_table(tmp_path):
    from PIL import Image
    from pypdf import PdfReader
    image=tmp_path/'pdf-image.png'
    Image.new('RGB',(100,100),'purple').save(image)
    path=tmp_path/'objects.pdf'
    call('pdf_create',engine='reportlab',output_path=str(path),blocks=[
        {'type':'table','rows':[['Field','Value'],['A long label that must wrap '*4,'42']]},
        {'type':'image','image_path':str(image),'width_cm':3,'caption':'Image caption'}])
    reader=PdfReader(path)
    assert '42' in reader.pages[0].extract_text()
    assert len(reader.pages[0].images)==1


def test_shared_ppt_chart_isolated_on_edit(tmp_path):
    from pptx import Presentation
    from copy import deepcopy
    path=deck(tmp_path,[element('chart','one',categories=['A'],series=[{'name':'S','values':[1]}])])
    prs=Presentation(path)
    slide=prs.slides[0]
    copy=deepcopy(slide.shapes[0]._element)
    copy.xpath('.//p:cNvPr')[0].set('id','3')
    copy.xpath('.//p:cNvPr')[0].set('name','two')
    slide.shapes._spTree.insert_element_before(copy,'p:extLst')
    prs.save(path)
    call('pptx_edit',path=str(path),slide_index=0,shape_name='one',operation='set_chart',
         chart_data={'categories':['A'],'series':[{'name':'S','values':[8]}]},in_place=True)
    shapes=Presentation(path).slides[0].shapes
    assert [list(s.chart.series[0].values) for s in shapes]==[[8],[1]]
    assert shapes[0].chart.part is not shapes[1].chart.part


def test_preview_region_page_mapping_and_pdf_layout(tmp_path,monkeypatch):
    from kiyori_office import render
    from PIL import Image
    from reportlab.pdfgen import canvas
    source=tmp_path/'two.pdf'
    pdf=canvas.Canvas(str(source))
    for text in ('First page','Second page'):
        pdf.drawString(72,700,text)
        pdf.showPage()
    pdf.save()
    commands=[]
    def render_page(command,output,*args,**kwargs):
        commands.append(command)
        Image.new('RGB',(100,200),'red').save(output)
    monkeypatch.setattr(render,'require_binary',lambda *args,**kwargs:'pdftoppm')
    monkeypatch.setattr(render,'require_poppler_data',lambda **kwargs:{'available':True})
    monkeypatch.setattr(render,'run_output_command',render_page)
    result=call('office_render_preview',path=str(source),pages='2',region={'left':0,'top':0,'width':0.5,'height':0.5},layout_report=True)
    assert result['artifacts'][0]['page']==2 and result['data']['remaining_pages']==[1]
    assert Image.open(result['artifacts'][0]['path']).size==(50,100)
    report=result['data']['layout_reports'][0]
    assert report['page']==2 and report['coordinate_origin']=='top_left'
    assert any(word['text']=='Second' for word in report['words'])
    assert commands[0][commands[0].index('-f')+1]=='2'


@pytest.mark.parametrize('omml', ['<!DOCTYPE x><m:oMath/>','<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><w:object xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/></m:oMath>'])
def test_unsafe_formula_never_published(tmp_path,omml):
    target=tmp_path/'bad.docx'
    result=protocol.run('docx_create',{'output_path':str(target),'spec':{'blocks':[{'type':'formula','omml':omml}]}})
    assert not result['ok'] and result['error']['code']=='E_INPUT_SCHEMA'
    assert not target.exists()


def test_validator_reports_dangling_animation(tmp_path):
    from pptx import Presentation
    path=deck(tmp_path)
    call('pptx_edit',path=str(path),slide_index=0,operation='set_animations',animations=[{'shape_name':'title'}],in_place=True)
    prs=Presentation(path)
    prs.slides[0]._element.xpath('.//p:spTgt')[0].set('spid','999')
    prs.save(path)
    result=protocol.run('office_validate',{'path':str(path),'strict':True})
    assert not result['ok']
    assert any(issue['code']=='ANIMATION_TARGET_MISSING' for issue in result['data']['issues'])
