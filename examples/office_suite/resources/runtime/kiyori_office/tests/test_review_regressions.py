"""2026-09-11 测评回归：核验实际保存部件、错误原子性与可读诊断。"""

from io import BytesIO
from zipfile import ZipFile

import pytest

from kiyori_office import protocol


def run(command, **args):
    result = protocol.run(command, args)
    assert result['ok'], result
    return result


def test_merge_preserves_images_math_and_independent_chart_workbooks(tmp_path):
    from docx import Document
    from PIL import Image
    from openpyxl import load_workbook
    from docx.oxml.ns import qn

    picture = tmp_path / 'plot.png'
    Image.new('RGB', (60, 40), 'green').save(picture)
    sources = [tmp_path / 'first.docx', tmp_path / 'second.docx']
    omml = '<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x=1</m:t></m:r></m:oMath>'
    for index, source in enumerate(sources):
        run('docx_create', output_path=str(source), spec={'blocks': [
            {'type': 'paragraph', 'text': 'Document %d' % index},
            {'type': 'image', 'image_path': str(picture), 'width_cm': 2},
            {'type': 'formula', 'omml': omml},
            {'type': 'chart', 'categories': ['A', 'B'],
             'series': [{'name': 'Series', 'values': [index + 1, index + 5]}]},
        ]})
    original_bytes = [source.read_bytes() for source in sources]
    output = tmp_path / 'merged.docx'
    run('docx_merge', paths=[str(source) for source in sources], output_path=str(output))
    assert [source.read_bytes() for source in sources] == original_bytes
    document = Document(output)
    assert len(document.element.body.xpath('.//m:oMath')) == 2
    drawings = document.element.body.xpath('.//wp:docPr')
    assert len({node.get('id') for node in drawings}) == len(drawings)
    charts = document.element.body.xpath('.//c:chart')
    assert len(charts) == 2
    workbooks = []
    for chart in charts:
        part = document.part.rels[chart.get(qn('r:id'))].target_part
        embedded = next(rel.target_part for rel in part.rels.values() if rel.reltype.endswith('/package'))
        workbook = load_workbook(BytesIO(embedded.blob))
        workbooks.append(workbook.active['B2'].value)
        workbook.close()
    assert workbooks == [1, 2]
    run('office_validate', path=str(output), strict=True)
    with ZipFile(output) as archive:
        assert len(archive.namelist()) == len(set(archive.namelist()))


def test_print_setup_and_formula_cache_diagnostics(tmp_path):
    from openpyxl import Workbook, load_workbook
    source, output = tmp_path / 'report.xlsx', tmp_path / 'print.xlsx'
    workbook = Workbook()
    workbook.active.append(['Amount', 'Total'])
    workbook.active.append([12, '=A2*2'])
    workbook.save(source)
    run('xlsx_format', path=str(source), output_path=str(output), print_setup={
        'print_area': 'A1:Q35', 'orientation': 'landscape', 'paper_size': 'A4',
        'fit_to_width': 1, 'fit_to_height': 0,
    })
    workbook = load_workbook(output)
    sheet = workbook.active
    assert '$A$1:$Q$35' in str(sheet.print_area)
    assert sheet.page_setup.orientation == 'landscape'
    assert sheet.page_setup.paperSize == int(sheet.PAPERSIZE_A4)
    assert sheet.sheet_properties.pageSetUpPr.fitToPage
    assert sheet.page_setup.fitToWidth == 1 and sheet.page_setup.fitToHeight == 0
    assert sheet['B2'].value == '=A2*2'
    workbook.close()
    read = run('office_read', path=str(output), mode='full')
    assert '缓存未提供' in read['data']['text']
    assert '=A2*2=None' not in read['data']['text']
    read = run('xlsx_read', path=str(output))
    formula = read['data']['sheets'][0]['rows'][1][1]
    assert formula['cached_value'] is None and formula['cache_status'] == 'missing'


@pytest.mark.parametrize('settings', [
    {'fit_to_width': -1}, {'fit_to_height': True}, {'orientation': 'diagonal'},
    {'print_area': 'A:Q'}, {'print_area': 'B2:A1'}, {'unknown': 1},
])
def test_invalid_print_settings_do_not_change_source(tmp_path, settings):
    from openpyxl import Workbook
    source = tmp_path / 'source.xlsx'
    Workbook().save(source)
    original = source.read_bytes()
    result = protocol.run('xlsx_format', {'path': str(source), 'in_place': True, 'print_setup': settings})
    assert not result['ok'] and result['error']['code'] == 'E_INPUT_SCHEMA'
    assert source.read_bytes() == original


def test_unknown_watermark_font_is_an_input_error_and_preserves_source(tmp_path):
    from pypdf import PdfWriter
    source = tmp_path / 'source.pdf'
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.write(source)
    original = source.read_bytes()
    result = protocol.run('pdf_watermark', {'path': str(source), 'in_place': True,
                                          'text': '中文', 'cjk_font': 'Missing-Font-123'})
    assert not result['ok'] and result['error']['code'] == 'E_INPUT_SCHEMA'
    assert source.read_bytes() == original


def test_font_family_registration_resolves_fontconfig_path(tmp_path, monkeypatch):
    import reportlab
    from reportlab.pdfbase import pdfmetrics
    from pathlib import Path
    from kiyori_office.ops_pdf import ops
    # Vera 只验证真实 TTF 注册和嵌入链路，不把西文字体当作中文渲染证明。
    font = Path(reportlab.__file__).parent / 'fonts' / 'Vera.ttf'
    monkeypatch.setattr(ops, 'require_cjk_fonts', lambda **kwargs: {
        'fonts': [], 'fontconfig_fonts': [{'family': 'Fixture family', 'file': str(font)}]})
    assert ops._register_reportlab_cjk_font('Fixture family', purpose='test') == 'Fixture family'
    assert pdfmetrics.getFont('Fixture family').face.filename == str(font)


def test_cid_alias_is_canonical_for_creation_and_watermark(tmp_path):
    from pypdf import PdfReader
    source, output = tmp_path / 'source.pdf', tmp_path / 'watermarked.pdf'
    created = run('pdf_create', engine='reportlab', output_path=str(source), cjk_font='STSong',
                  blocks=[{'type': 'paragraph', 'text': '中文正文'}])
    assert created['data']['font'] == 'STSong-Light'
    run('pdf_watermark', path=str(source), output_path=str(output), cjk_font='STSong', text='测试水印')
    assert '测试水印' in PdfReader(output).pages[0].extract_text()
