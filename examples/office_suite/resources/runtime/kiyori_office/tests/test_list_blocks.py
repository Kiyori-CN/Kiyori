"""列表输入必须落成真实段落；软换行与段落边界保持各自语义。"""
import pytest
from docx import Document
from pptx import Presentation

from kiyori_office import protocol


@pytest.mark.parametrize('kind,style', [('bullet', 'List Bullet'), ('number', 'List Number')])
def test_docx_list_items_preserve_order_and_format(tmp_path, kind, style):
    output = tmp_path / 'list.docx'
    result = protocol.run('docx_create', {'output_path': str(output), 'spec': {'blocks': [
        {'type': 'paragraph', 'text': '前文'},
        {'type': kind, 'items': ['第一项', '第二项', '第三项'], 'format': {'space_after_pt': 7}},
        {'type': 'paragraph', 'text': '后文'},
    ]}})
    assert result['ok'], result
    paragraphs = Document(output).paragraphs
    assert [p.text for p in paragraphs] == ['前文', '第一项', '第二项', '第三项', '后文']
    for p in paragraphs[1:4]:
        assert p.style.name == style
        assert p.paragraph_format.space_after.pt == 7
    assert result['data']['paragraph_count'] == 5
    assert protocol.run('office_validate', {'path': str(output), 'strict': True})['ok']


@pytest.mark.parametrize('block', [
    {'type': 'bullet', 'items': []}, {'type': 'bullet', 'items': 'one'},
    {'type': 'bullet', 'items': [None]}, {'type': 'number', 'items': [1]},
    {'type': 'bullet', 'items': ['one'], 'text': 'two'},
    {'type': 'bullet', 'items': ['one'], 'runs': [{'text': 'two'}]},
    {'type': 'bullet'}, {'type': 'number', 'runs': []},
    {'type': 'bullet', 'itmes': ['one']}, {'type': 'paragraph', 'items': ['one']},
])
def test_invalid_list_or_misspelled_content_never_overwrites(tmp_path, block):
    output = tmp_path / 'original.docx'
    doc = Document(); doc.add_paragraph('原文'); doc.save(output)
    original = output.read_bytes()
    result = protocol.run('docx_create', {'output_path': str(output), 'overwrite': True, 'spec': {'blocks': [block]}})
    assert not result['ok'], result
    assert result['error']['code'] == 'E_INPUT_SCHEMA'
    assert output.read_bytes() == original


def test_docx_single_item_text_and_runs_remain_compatible(tmp_path):
    output = tmp_path / 'list.docx'
    result = protocol.run('docx_create', {'output_path': str(output), 'spec': {'blocks': [
        {'type': 'bullet', 'text': '单项'},
        {'type': 'number', 'runs': [{'text': '重点', 'bold': True}, {'text': '说明'}]},
    ]}})
    assert result['ok'], result
    paragraphs = Document(output).paragraphs
    assert [p.text for p in paragraphs] == ['单项', '重点说明']
    assert paragraphs[1].runs[0].bold is True


def test_pptx_element_soft_break_and_explicit_paragraphs(tmp_path):
    output = tmp_path / 'lines.pptx'
    geometry = {'type': 'text', 'left_cm': 1, 'top_cm': 1, 'width_cm': 10, 'height_cm': 4}
    result = protocol.run('pptx_create', {'output_path': str(output), 'layout_index': 6, 'slides': [{'elements': [
        {**geometry, 'text': '第一行\n第二行'},
        {**geometry, 'top_cm': 6, 'paragraphs': ['第一段', '第二段']},
    ]}]})
    assert result['ok'], result
    shapes = Presentation(output).slides[0].shapes
    assert len(shapes[0].text_frame.paragraphs) == 1
    assert shapes[0].text == '第一行\v第二行'
    assert len(shapes[1].text_frame.paragraphs) == 2
    assert shapes[1].text == '第一段\n第二段'
