"""日常表格与文档工作流：回读真实产物，验证保真及错误路径。"""
import json
import zipfile
from unittest.mock import patch

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from kiyori_office import protocol
from kiyori_office.ops_xlsx.formula import normalize_formula, spill_functions


@pytest.mark.parametrize('formula', ['=FILTER(A1:A3,A1:A3>0)', '=_xlfn._xlws.FILTER(A1:A3,A1:A3>0)'])
def test_value_formula_cannot_bypass_spill_policy(tmp_path, formula):
    output = tmp_path / 'bad.xlsx'
    result = protocol.run('xlsx_write', {'cells': [{'cell':'A1','value':formula}], 'output_path':str(output)})
    assert result['error']['code'] == 'E_INPUT_SCHEMA'
    assert not output.exists()


def test_formula_literals_are_not_parsed_as_functions():
    formula = '=IF(A1="FILTER(",CONCAT("IFS(",B1),"UNIQUE(")'
    assert normalize_formula(formula) == '=IF(A1="FILTER(",_xlfn.CONCAT("IFS(",B1),"UNIQUE(")'
    assert spill_functions(formula) == []
    assert spill_functions('=_xlfn._xlws.FILTER(A1:A3,A1:A3>0)') == ['FILTER']


def test_range_format_preserves_other_cells_and_font_attributes(tmp_path):
    source, output = tmp_path / 'source.xlsx', tmp_path / 'format.xlsx'
    book = Workbook()
    book.active.append(['Quarter','Revenue'])
    book.active.append(['Q1',10])
    book.active['A1'].font = Font(name='Arial', bold=True, color='FF112233')
    book.save(source)
    result = protocol.run('xlsx_format', dict(path=str(source), output_path=str(output), range='A1:B1', font={'size':18}, fill={'color':'FFEEEEEE'}))
    assert result['ok'], result
    edited = load_workbook(output)
    assert edited.active['A1'].font.bold and edited.active['A1'].font.name == 'Arial'
    assert edited.active['A1'].font.color.rgb == 'FF112233'
    assert edited.active['B1'].font.sz == 18
    assert edited.active['A2'].font.sz != 18
    assert edited.active['A2'].fill.patternType is None


@pytest.mark.parametrize('kind', ['column','bar','line','pie','scatter'])
def test_native_charts_survive_save_and_reference_real_cells(tmp_path, kind):
    source, output = tmp_path / 'data.xlsx', tmp_path / (kind+'.xlsx')
    book = Workbook()
    book.active.title = '季度数据'
    for row in [['Quarter','Revenue'],[1,10],[2,20],[3,15]]:
        book.active.append(row)
    book.save(source)
    result = protocol.run('xlsx_chart', dict(path=str(source), output_path=str(output), data_range='A1:B4', chart_type=kind, title='营收', anchor='D2'))
    assert result['ok'], result
    edited = load_workbook(output)
    assert edited.active['B4'].value == 15
    assert len(edited.active._charts) == 1
    assert len(edited.active._charts[0].series) == 1
    with zipfile.ZipFile(output) as archive:
        xml = archive.read('xl/charts/chart1.xml').decode('utf-8')
        assert '$B$2:$B$4' in xml and '$A$2:$A$4' in xml
    assert protocol.run('office_validate', dict(path=str(output), strict=True))['ok']


def test_pie_rejects_multiple_series_without_output(tmp_path):
    source, output = tmp_path/'data.xlsx', tmp_path/'pie.xlsx'
    book = Workbook()
    book.active.append(['Quarter','Revenue','Cost'])
    book.active.append(['Q1',10,8])
    book.save(source)
    result = protocol.run('xlsx_chart',dict(path=str(source),output_path=str(output),data_range='A1:C2',chart_type='pie'))
    assert result['error']['code'] == 'E_INPUT_SCHEMA'
    assert not output.exists()


def test_macro_extension_change_is_rejected_before_open(tmp_path):
    source, output = tmp_path/'macro.xlsm', tmp_path/'plain.xlsx'
    source.write_bytes(b'unchanged macro template')
    result = protocol.run('xlsx_write',dict(path=str(source),output_path=str(output),rows=[[1]]))
    assert result['error']['code'] == 'E_INPUT_SCHEMA'
    assert not output.exists() and source.read_bytes() == b'unchanged macro template'


def test_merged_word_cells_are_replaced_once(tmp_path):
    source, output = tmp_path/'table.docx', tmp_path/'edited.docx'
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0,0).merge(table.cell(0,1)).text = 'A'
    doc.save(source)
    result = protocol.run('docx_find_replace',dict(path=str(source),output_path=str(output),find='A',replace='AA'))
    assert result['ok'], result
    assert result['data']['replacements'] == 1
    assert Document(output).tables[0].cell(0,0).text == 'AA'


def test_failed_recalc_does_not_publish_or_replace(tmp_path):
    from kiyori_office.ops_xlsx import recalc
    source, output = tmp_path/'source.xlsx', tmp_path/'important.xlsx'
    book = Workbook()
    book.active['A1'] = '=1+1'
    book.save(source)
    output.write_bytes(b'original')
    def fake_soffice(soffice, profile, source, target_dir):
        (target_dir / source.name).write_bytes(source.read_bytes())
        return {'exit_code':0}
    with patch.object(recalc, 'require_libreoffice', return_value={'path':'soffice'}), patch.object(recalc,'_run_soffice',side_effect=fake_soffice):
        result = protocol.run('xlsx_recalc',dict(path=str(source),output_path=str(output),overwrite=True))
    assert result['error']['code'] == 'E_VALIDATION_FAILED'
    assert result['data']['missing_cache_count'] == 1
    assert output.read_bytes() == b'original'


def test_large_structured_output_is_bounded_and_recoverable(tmp_path):
    source = tmp_path/'many.xlsx'
    book = Workbook()
    for i in range(100):
        book.active.append([i,'详细文本'*20])
    book.save(source)
    result = protocol.run('office_read',dict(path=str(source),mode='full',max_chars=1000))
    assert result['ok'] and result['truncated']
    from pathlib import Path
    full = json.loads(Path(result['full_output_path']).read_text(encoding='utf-8'))
    assert full
    assert len(json.dumps(result['data'],ensure_ascii=False)) < 1300


@pytest.mark.parametrize('failure', ['exit','empty','timeout'])
def test_external_engine_failure_preserves_existing_document(tmp_path, failure):
    import subprocess
    from pathlib import Path
    from kiyori_office.engines import run_output_command
    output = tmp_path/'important.pdf'
    output.write_bytes(b'original')
    def engine(command, **kwargs):
        if failure == 'timeout':
            Path(command[-1]).write_bytes(b'partial')
            raise subprocess.TimeoutExpired(command,1)
        if failure == 'exit':
            Path(command[-1]).write_bytes(b'partial')
        return subprocess.CompletedProcess(command,1 if failure=='exit' else 0,'','engine failed')
    with patch('kiyori_office.engines.subprocess.run',side_effect=engine), pytest.raises(protocol.OfficeError) as exc:
        run_output_command(['engine',str(output)],output,-1,timeout=1,purpose='test')
    assert exc.value.code == ('E_TIMEOUT' if failure=='timeout' else 'E_ENGINE_FAILED')
    assert output.read_bytes() == b'original'
    assert list(tmp_path.glob('.tmp-*')) == []


def test_sparse_workbook_info_does_not_expand_used_rectangle(tmp_path):
    from openpyxl.worksheet.worksheet import Worksheet
    source = tmp_path/'sparse.xlsx'
    book = Workbook()
    book.active['A1'] = '=1+1'
    book.active['XFD1048576'].font = Font(bold=True)
    book.save(source)
    with patch.object(Worksheet,'iter_rows',side_effect=AssertionError('dense iteration')):
        result = protocol.run('xlsx_info',dict(path=str(source)))
    assert result['ok'],result
    assert result['data']['formula_cells'] == 1
