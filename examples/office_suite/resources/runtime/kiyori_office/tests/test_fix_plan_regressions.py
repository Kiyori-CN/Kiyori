"""修复方案的可观察回归：真实 PDF/OOXML 回读，外部引擎仅验证调用契约。"""
import subprocess
from contextlib import nullcontext
from pathlib import Path
from types import SimpleNamespace

import pytest

from kiyori_office import protocol


def make_form(path, radio=False):
    from reportlab.pdfgen.canvas import Canvas
    canvas = Canvas(str(path))
    if radio:
        for index, value in enumerate(('Yes', 'No')):
            canvas.acroForm.radio(name='answer', value=value, x=40, y=700-index*40,
                                 selected=index == 0)
    else:
        canvas.acroForm.checkbox(name='agree', x=40, y=700, checked=False)
    canvas.showPage()
    canvas.save()


@pytest.mark.parametrize('value,state', [('No', '/No'), ('Yes', '/Yes'), ('/No', '/No')])
def test_radio_uses_real_export_state_and_disables_siblings(tmp_path, value, state):
    from pypdf import PdfReader
    from pypdf.generic import NameObject
    source, output = tmp_path/'source.pdf', tmp_path/'filled.pdf'
    make_form(source, radio=True)
    result = protocol.run('pdf_form_fill', {'path': str(source), 'output_path': str(output),
                                          'values': {'answer': value}})
    assert result['ok'], result
    reader = PdfReader(output)
    assert reader.get_fields()['answer']['/V'] == state
    widgets = [ref.get_object() for ref in reader.pages[0]['/Annots']]
    assert [w['/AS'] for w in widgets] == ([state, '/Off'] if state == '/Yes' else ['/Off', state])
    assert all(isinstance(w['/AS'], NameObject) for w in widgets)


def test_ambiguous_radio_bool_preserves_existing_output(tmp_path):
    source, output = tmp_path/'source.pdf', tmp_path/'filled.pdf'
    make_form(source, radio=True)
    output.write_bytes(b'existing output')
    result = protocol.run('pdf_form_fill', {'path': str(source), 'output_path': str(output),
        'overwrite': True, 'values': {'answer': True}})
    assert not result['ok'] and result['error']['code'] == 'E_INPUT_SCHEMA'
    assert 'available_states' in result['error']['detail']
    assert output.read_bytes() == b'existing output'


def test_nested_checkbox_and_unknown_fields_report_actual_writes(tmp_path):
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import ArrayObject, DictionaryObject, NameObject, TextStringObject
    source, nested, output = [tmp_path/name for name in ('source.pdf', 'nested.pdf', 'filled.pdf')]
    make_form(source)
    writer = PdfWriter(clone_from=source)
    child_ref = writer._root_object['/AcroForm']['/Fields'][0]
    parent = DictionaryObject({NameObject('/T'): TextStringObject('section'),
                               NameObject('/Kids'): ArrayObject([child_ref])})
    parent_ref = writer._add_object(parent)
    child_ref.get_object()[NameObject('/Parent')] = parent_ref
    writer._root_object['/AcroForm'][NameObject('/Fields')] = ArrayObject([parent_ref])
    writer.write(nested)
    result = protocol.run('pdf_form_fill', {'path': str(nested), 'output_path': str(output),
        'strict': False, 'values': {'section.agree': True, 'missing': True}})
    assert result['ok'], result
    assert result['data']['filled'] == ['section.agree']
    assert result['data']['unknown'] == ['missing']
    assert PdfReader(output).get_fields()['section.agree']['/V'] == '/Yes'
    listed = protocol.run('pdf_form_list', {'path': str(nested)})
    field = next(f for f in listed['data']['fields'] if f['name'] == 'section.agree')
    assert set(field['available_states']) == {'/Yes', '/Off'}


@pytest.mark.parametrize('broken', ['pushbutton', 'no_widget'])
def test_nonfillable_button_does_not_claim_success(tmp_path, broken):
    from pypdf import PdfWriter
    from pypdf.generic import NameObject, NumberObject, ArrayObject
    source, changed, output = [tmp_path/name for name in ('source.pdf', 'changed.pdf', 'out.pdf')]
    make_form(source)
    writer = PdfWriter(clone_from=source)
    if broken == 'pushbutton':
        writer._root_object['/AcroForm']['/Fields'][0].get_object()[NameObject('/Ff')] = NumberObject(65536)
    else:
        writer.pages[0][NameObject('/Annots')] = ArrayObject()
    writer.write(changed)
    result = protocol.run('pdf_form_fill', {'path': str(changed), 'output_path': str(output),
                                           'values': {'agree': True}})
    assert not result['ok'] and result['error']['code'] == 'E_INPUT_SCHEMA', result
    assert not output.exists()


@pytest.mark.parametrize('initial,retry,code,calls', [
    ([[['A']]], [], None, 1), ([], [], None, 1),
    ([[['', None]]], [[['季度', '销售额'], ['一季度', '120.00']]], 'PDF_TABLE_STRATEGY_FALLBACK', 2),
    ([[['', None]]], [[['']]], 'PDF_TABLE_TEXT_EMPTY', 2),
    ([[['', None]]], [], 'PDF_TABLE_TEXT_EMPTY', 2),
])
def test_table_strategy_is_conditional_and_never_silently_empty(tmp_path, monkeypatch, initial, retry, code, calls):
    import pdfplumber
    from kiyori_office.ops_pdf.ops import _extract_tables
    observed = []
    def extract(settings=None):
        observed.append(settings)
        return initial if settings is None else retry
    monkeypatch.setattr(pdfplumber, 'open', lambda path: nullcontext(
        SimpleNamespace(pages=[SimpleNamespace(extract_tables=extract)])))
    result = _extract_tables(tmp_path/'input.pdf', [1], tmp_path, {})
    assert len(observed) == calls
    assert [w['code'] for w in result['warnings']] == ([code] if code else [])
    if code == 'PDF_TABLE_STRATEGY_FALLBACK':
        assert result['data']['tables'][0]['rows'] == retry[0]
        assert observed[1] == {'vertical_strategy': 'text', 'horizontal_strategy': 'text'}


@pytest.mark.parametrize('suffix,target,writer_import', [
    ('.pdf', 'docx', True), ('.PDF', 'odt', True), ('.pdf', 'html', True),
    ('.docx', 'pdf', False), ('.xlsx', 'pdf', False), ('.pdf', 'pptx', False),
])
def test_libreoffice_conversion_filter_and_engine(tmp_path, monkeypatch, suffix, target, writer_import):
    from kiyori_office import convert
    source, output = tmp_path/('source'+suffix), tmp_path/('output.'+target)
    source.write_bytes(b'command contract fixture')
    monkeypatch.setattr(convert, 'require_libreoffice_for_source', lambda *a, **kw: {'path': 'soffice', 'version': 'test'})
    def require_writer(**kwargs):
        assert writer_import and kwargs['component'] == 'writer'
        return {'path': 'soffice', 'version': 'test'}
    monkeypatch.setattr(convert, 'require_libreoffice', require_writer)
    def engine(command, **kwargs):
        assert ('--infilter=writer_pdf_import' in command) is writer_import
        if writer_import:
            assert command.index('--infilter=writer_pdf_import') < command.index('--convert-to')
        destination = Path(command[command.index('--outdir')+1])/(source.stem+'.'+target)
        destination.write_bytes(b'converted by test double')
        return subprocess.CompletedProcess(command, 0, '', '')
    monkeypatch.setattr(convert.subprocess, 'run', engine)
    result = protocol.run('office_convert', {'from_path': str(source), 'to_format': target,
        'engine': 'libreoffice', 'output_path': str(output)})
    assert result['ok'], result
    assert result['engine'] == result['data']['engine'] == 'libreoffice'
    assert output.read_bytes() == b'converted by test double'


def test_weasyprint_plan_and_result_engine(tmp_path, monkeypatch):
    from kiyori_office import env, engines
    from kiyori_office.ops_pdf import ops
    from pypdf import PdfWriter
    assert 'weasyprint' in env.TIER_COMPONENTS[2]
    assert any('weasyprint' in command for command in env.build_plan(2)['commands'])
    assert any('weasyprint' in command for command in env.build_plan(2, ['weasyprint'])['commands'])
    monkeypatch.setattr(ops, 'require_binary', lambda *a, **kw: 'weasyprint')
    monkeypatch.setattr(ops, 'require_cjk_font_files', lambda **kwargs: {})
    monkeypatch.setattr(ops, 'engine_version', lambda engine: 'test-version' if engine == 'weasyprint' else None)
    def engine(command, output, *a, **kw):
        writer = PdfWriter(); writer.add_blank_page(width=200, height=200); writer.write(output)
    monkeypatch.setattr(engines, 'run_output_command', engine)
    source, output = tmp_path/'source.html', tmp_path/'out.pdf'
    source.write_text('<p>text</p>', encoding='utf-8')
    result = protocol.run('pdf_create', {'engine': 'weasyprint', 'source_path': str(source), 'output_path': str(output)})
    assert result['ok'], result
    assert result['engine'] == result['data']['engine'] == 'weasyprint'
    assert result['engine_version'] == 'test-version'


@pytest.mark.parametrize('mode', ['filled', 'strict', 'loose', 'no_body'])
def test_template_notes_are_filled_without_creating_notes(tmp_path, mode):
    from pptx import Presentation
    source, output = tmp_path/'source.pptx', tmp_path/'out.pptx'
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    prs.slides.add_slide(prs.slide_layouts[6])
    frame = slide.notes_slide.notes_text_frame
    frame.text = '备注：'
    frame.paragraphs[0].add_run().text = '{{no'
    frame.paragraphs[0].add_run().text = 'te}}'
    if mode == 'no_body':
        body = slide.notes_slide.notes_placeholder._element
        body.getparent().remove(body)
    prs.save(source)
    result = protocol.run('pptx_template_fill', {'path': str(source), 'output_path': str(output),
        'variables': {'note': '替换成功'} if mode == 'filled' else {}, 'strict': mode != 'loose'})
    if mode == 'strict':
        assert not result['ok'] and result['error']['code'] == 'E_TEMPLATE_VAR_MISSING'
        assert not output.exists()
        return
    assert result['ok'], result
    restored = Presentation(output)
    assert not restored.slides[1].has_notes_slide
    if mode == 'no_body':
        assert restored.slides[0].notes_slide.notes_text_frame is None
    else:
        assert restored.slides[0].notes_slide.notes_text_frame.text == ('备注：替换成功' if mode == 'filled' else '备注：{{note}}')
        assert result['data']['missing_variables'] == (['note'] if mode == 'loose' else [])


@pytest.mark.parametrize('explicit', [False, True])
def test_shape_alignment_and_compact_table_rows_roundtrip(tmp_path, explicit):
    from pptx import Presentation
    from pptx.enum.text import MSO_ANCHOR
    box = {'left_cm': 1, 'top_cm': 1, 'width_cm': 10, 'height_cm': 10}
    shape = {**box, 'type': 'shape', 'text': 'container'}
    table = {**box, 'type': 'table', 'rows': [['a'], ['b']]}
    if explicit:
        shape['vertical_alignment'] = 'top'
        table['table_style'] = {'row_heights_cm': [3, 2]}
    output = tmp_path/'out.pptx'
    result = protocol.run('pptx_create', {'output_path': str(output), 'slides': [
        {'layout_index': 6, 'elements': [shape, table, {**box, 'type': 'text', 'text': 'plain'}]}]})
    assert result['ok'], result
    shapes = Presentation(output).slides[0].shapes
    assert shapes[0].text_frame.vertical_anchor == (MSO_ANCHOR.TOP if explicit else MSO_ANCHOR.MIDDLE)
    assert [row.height.cm for row in shapes[1].table.rows] == ([3, 2] if explicit else [0.8, 0.8])
    assert shapes[2].text_frame.vertical_anchor == MSO_ANCHOR.TOP


def test_unknown_schema_field_lists_fields_at_its_level():
    result = protocol.run('pptx_create', {'slides': [], 'transition': {}})
    assert not result['ok'] and result['error']['code'] == 'E_INPUT_SCHEMA'
    issue = next(i for i in result['data']['issues'] if i['field'] == 'args.transition')
    assert 'slides' in issue['expected'] and 'transition' not in issue['expected']
    assert '该层级可用字段' in result['error']['message']


def test_compact_table_actual_height_cannot_escape_slide(tmp_path):
    output = tmp_path/'out.pptx'
    result = protocol.run('pptx_create', {'output_path': str(output), 'slides': [
        {'layout_index': 6, 'elements': [{'type': 'table', 'left_cm': 1, 'top_cm': 1,
            'width_cm': 10, 'height_cm': 10, 'rows': [['row'] for _ in range(30)]}]}]})
    assert not result['ok'] and result['error']['code'] == 'E_INPUT_SCHEMA'
    assert not output.exists()


def test_real_pdf_empty_grid_recovers_text_table_with_warning(tmp_path):
    from reportlab.pdfgen.canvas import Canvas
    import pdfplumber
    path = tmp_path/'offset-grid.pdf'
    canvas = Canvas(str(path))
    # 模拟空网格与文本不重合的文件，直接核验真实 pdfplumber 两种策略的差异。
    for x in (50, 150, 250, 350):
        canvas.line(x, 100, x, 250)
    for y in (100, 150, 200, 250):
        canvas.line(50, y, 350, y)
    for row in range(4):
        for column, text in enumerate(('Quarter', 'Sales', 'Profit')):
            canvas.drawString(50 + column*100, 700-row*30, text if row == 0 else str(row*100+column))
    canvas.showPage(); canvas.save()
    with pdfplumber.open(path) as pdf:
        assert all(not cell for table in pdf.pages[0].extract_tables() for row in table for cell in row)
    result = protocol.run('pdf_extract', {'path': str(path), 'mode': 'tables'})
    assert result['ok'], result
    assert 'PDF_TABLE_STRATEGY_FALLBACK' in [w['code'] for w in result['warnings']]
    cells = [cell for table in result['data']['tables'] for row in table['rows'] for cell in row]
    assert 'Quarter' in cells and '301' in cells
