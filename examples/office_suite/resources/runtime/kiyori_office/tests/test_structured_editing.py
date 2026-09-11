"""论文与演示文稿局部编辑：使用真实文件证明结构和未命中内容保真。"""
from copy import deepcopy

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches

from kiyori_office import protocol
from kiyori_office.ops_docx.layout import append_field
from kiyori_office.validate import validate_document


def run(command, source, output, **args):
    return protocol.run(command, dict(path=str(source), output_path=str(output), **args))


def test_docx_replace_rejects_field_without_touching_source(tmp_path):
    source = tmp_path / 'paper.docx'
    doc = Document()
    p = doc.add_paragraph('见图 ')
    append_field(p, 'REF figure_one', '1')
    doc.save(source)
    original = source.read_bytes()
    result = run('docx_edit', source, source, operation='replace', anchor={'index': 0}, text='新正文', in_place=True)
    assert not result['ok'], result
    assert result['error']['code'] == 'E_FORMAT_UNSUPPORTED'
    assert source.read_bytes() == original


@pytest.mark.parametrize('operation', ['insert_before', 'insert_after'])
def test_docx_insert_inherits_format_without_duplicating_section(tmp_path, operation):
    source, output = tmp_path / 'paper.docx', tmp_path / 'edited.docx'
    doc = Document()
    p = doc.add_paragraph('原段', style='Heading 1')
    p.runs[0].font.name = 'Arial'
    p.runs[0].font.size = Pt(18)
    p._p.get_or_add_pPr().append(deepcopy(doc.sections[0]._sectPr))
    doc.save(source)
    original_sections = len(doc.sections)
    result = run('docx_edit', source, output, operation=operation, anchor={'index': 0}, text='新段')
    assert result['ok'], result
    reopened = Document(output)
    assert len(reopened.sections) == original_sections
    inserted = next(p for p in reopened.paragraphs if p.text == '新段')
    assert inserted.style.name == 'Heading 1'
    assert inserted.runs[0].font.size.pt == 18
    assert inserted.runs[0].font.name == 'Arial'


def test_docx_template_is_one_pass_and_visits_nested_and_even_stories(tmp_path):
    source, output = tmp_path / 'paper.docx', tmp_path / 'filled.docx'
    doc = Document()
    doc.add_paragraph('{{a}} {{b}} {{a}}')
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.add_table(rows=1, cols=1).cell(0, 0).text = '{{nested}}'
    doc.sections[0].different_first_page_header_footer = True
    doc.sections[0].first_page_header.paragraphs[0].text = '{{first}}'
    doc.settings.odd_and_even_pages_header_footer = True
    doc.sections[0].even_page_footer.paragraphs[0].text = '{{even}}'
    doc.save(source)
    result = run('docx_from_template', source, output, variables={'a': '{{b}}', 'b': 'B', 'nested': 'N', 'first': 'F', 'even': 'E'})
    assert result['ok'], result
    reopened = Document(output)
    assert reopened.paragraphs[0].text == '{{b}} B {{b}}'
    assert reopened.tables[0].cell(0, 0).tables[0].cell(0, 0).text == 'N'
    assert reopened.sections[0].first_page_header.paragraphs[0].text == 'F'
    assert reopened.sections[0].even_page_footer.paragraphs[0].text == 'E'


def test_docx_find_replace_does_not_edit_field_cache(tmp_path):
    source, output = tmp_path / 'paper.docx', tmp_path / 'edited.docx'
    doc = Document()
    p = doc.add_paragraph('旧 ')
    for kind in ('begin', 'separate'):
        marker = OxmlElement('w:fldChar'); marker.set(qn('w:fldCharType'), kind)
        p.add_run()._r.append(marker)
    p.add_run('旧')
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    p.add_run()._r.append(end)
    p.add_run(' 旧')
    doc.save(source)
    result = run('docx_find_replace', source, output, find='旧', replace='新')
    assert result['ok'], result
    assert Document(output).paragraphs[0].text == '新 旧 新'
    assert result['data']['replacements'] == 2


def test_docx_validator_reports_bookmark_and_anchor_errors(tmp_path):
    source = tmp_path / 'paper.docx'
    doc = Document()
    p = doc.add_paragraph('正文')
    for name in ('same', 'same'):
        marker = OxmlElement('w:bookmarkStart'); marker.set(qn('w:id'), '8'); marker.set(qn('w:name'), name)
        p._p.append(marker)
    link = OxmlElement('w:hyperlink'); link.set(qn('w:anchor'), 'missing')
    p._p.append(link)
    doc.save(source)
    codes = {i['code'] for i in validate_document(source)['issues']}
    assert {'BOOKMARK_DUPLICATE', 'BOOKMARK_UNPAIRED', 'HYPERLINK_ANCHOR_MISSING'} <= codes


def test_pptx_group_template_and_soft_break_boundaries(tmp_path):
    source, output = tmp_path / 'deck.pptx', tmp_path / 'filled.pptx'
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    nested = group.shapes.add_group_shape()
    box = nested.shapes.add_textbox(0, 0, Inches(2), Inches(1))
    box.text = '{{title}}'
    boundary = slide.shapes.add_textbox(0, Inches(2), Inches(2), Inches(1))
    p = boundary.text_frame.paragraphs[0]
    p.add_run().text = '{{ti'
    p.add_line_break()
    p.add_run().text = 'tle}}'
    prs.save(source)
    result = run('pptx_template_fill', source, output, variables={'title': '结论'})
    assert result['ok'], result
    reopened = Presentation(output)
    assert reopened.slides[0].shapes[0].shapes[0].shapes[0].text == '结论'
    assert reopened.slides[0].shapes[1].text == '{{ti\vtle}}'


def test_pptx_delete_rejects_inbound_link_and_preserves_source(tmp_path):
    source = tmp_path / 'deck.pptx'
    prs = Presentation()
    a = prs.slides.add_slide(prs.slide_layouts[1])
    b = prs.slides.add_slide(prs.slide_layouts[1])
    a.shapes.title.click_action.target_slide = b
    prs.save(source)
    original = source.read_bytes()
    result = run('pptx_slide', source, source, operation='delete', index=1, in_place=True)
    assert not result['ok'], result
    assert result['error']['code'] == 'E_FORMAT_UNSUPPORTED'
    assert 'slide' in result['error']['detail']
    assert source.read_bytes() == original


def test_pptx_edit_rejects_ambiguous_shape_name(tmp_path):
    source, output = tmp_path / 'deck.pptx', tmp_path / 'edited.pptx'
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for _ in range(2):
        s.shapes.add_textbox(0, 0, Inches(1), Inches(1)).name = 'title'
    prs.save(source)
    result = run('pptx_edit', source, output, slide_index=0, shape_name='title', text='修改')
    assert not result['ok'], result
    assert result['error']['code'] == 'E_ANCHOR_NOT_FOUND'
    assert not output.exists()


def test_docx_template_missing_nested_variable_does_not_publish(tmp_path):
    source, output = tmp_path / 'template.docx', tmp_path / 'filled.docx'
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = '{{missing}}'
    doc.save(source)
    result = run('docx_from_template', source, output, variables={})
    assert not result['ok']
    # D6：模板变量缺失是独立语义，不再复用「锚点未找到」的错误码。
    assert result['error']['code'] == 'E_TEMPLATE_VAR_MISSING'
    assert not output.exists()


def test_docx_template_keeps_header_parts_and_untouched_xml(tmp_path):
    from zipfile import ZipFile
    source, output = tmp_path / 'template.docx', tmp_path / 'filled.docx'
    doc = Document()
    p = doc.add_paragraph('未命中')
    p.add_run('也不合并')
    untouched = p._p.xml
    doc.add_paragraph('{{value}}')
    doc.save(source)
    result = run('docx_from_template', source, output, variables={'value': ''})
    assert result['ok'], result
    assert result['data']['replacements'] == 1
    assert Document(output).paragraphs[0]._p.xml == untouched
    with ZipFile(source) as before, ZipFile(output) as after:
        assert set(before.namelist()) == set(after.namelist())


@pytest.mark.parametrize('strict', [True, False])
def test_docx_protected_placeholder_is_never_reported_as_filled(tmp_path, strict):
    source, output = tmp_path / 'template.docx', tmp_path / 'filled.docx'
    doc = Document()
    p = doc.add_paragraph()
    append_field(p, 'REF test', '{{value}}')
    doc.save(source)
    result = run('docx_from_template', source, output, variables={'value': 'new'}, strict=strict)
    if strict:
        assert not result['ok'] and result['error']['code'] == 'E_FORMAT_UNSUPPORTED'
        assert not output.exists()
    else:
        assert result['ok'] and result['warnings']
        assert result['data']['replacements'] == 0
        assert result['data']['used_variables'] == []
        assert Document(output).paragraphs[0].text == '{{value}}'


def test_docx_find_replace_protects_multi_paragraph_field(tmp_path):
    source, output = tmp_path / 'paper.docx', tmp_path / 'edited.docx'
    doc = Document()
    start = OxmlElement('w:fldChar'); start.set(qn('w:fldCharType'), 'begin')
    doc.add_paragraph().add_run()._r.append(start)
    doc.add_paragraph('旧')
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    doc.add_paragraph().add_run()._r.append(end)
    doc.add_paragraph('旧')
    doc.save(source)
    result = run('docx_find_replace', source, output, find='旧', replace='新')
    assert result['ok'], result
    assert [p.text for p in Document(output).paragraphs] == ['', '旧', '', '新']


@pytest.mark.parametrize('operation,index,allowed', [('insert_before', 0, True), ('insert_after', 0, False),
                                                   ('insert_before', 1, False), ('insert_after', 1, True)])
def test_docx_insertion_cannot_enter_multi_paragraph_field(tmp_path, operation, index, allowed):
    source, output = tmp_path / 'paper.docx', tmp_path / 'edited.docx'
    doc = Document()
    for kind in ('begin', 'end'):
        marker = OxmlElement('w:fldChar'); marker.set(qn('w:fldCharType'), kind)
        doc.add_paragraph().add_run()._r.append(marker)
    doc.save(source)
    result = run('docx_edit', source, output, operation=operation, anchor={'index': index}, text='新段')
    assert result['ok'] == allowed, result
    if not allowed:
        assert result['error']['code'] == 'E_FORMAT_UNSUPPORTED'
        assert not output.exists()


def test_docx_valid_bookmark_and_internal_link_pass(tmp_path):
    source = tmp_path / 'paper.docx'
    doc = Document()
    p = doc.add_paragraph('图1')
    start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), '8'); start.set(qn('w:name'), 'figure_one')
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), '8')
    p._p.insert(0, start)
    p._p.append(end)
    link = OxmlElement('w:hyperlink'); link.set(qn('w:anchor'), 'figure_one')
    doc.add_paragraph('见图1')._p.append(link)
    doc.save(source)
    assert validate_document(source, strict=True)['issue_count'] == 0


def test_pptx_add_delete_move_and_duplicate_keep_navigation(tmp_path):
    source, added, deleted, moved, copied = [tmp_path / (name + '.pptx') for name in ('source', 'added', 'deleted', 'moved', 'copied')]
    prs = Presentation()
    a = prs.slides.add_slide(prs.slide_layouts[1])
    b = prs.slides.add_slide(prs.slide_layouts[1])
    a.shapes.title.text, b.shapes.title.text = 'A', 'B'
    a.shapes.title.click_action.target_slide = b
    prs.save(source)
    assert run('pptx_slide', source, added, operation='add', layout_index=6)['ok']
    assert run('pptx_slide', added, deleted, operation='delete', index=2)['ok']
    assert run('pptx_slide', deleted, moved, operation='move', index=1, target_index=0)['ok']
    assert run('pptx_slide', moved, copied, operation='duplicate', index=1)['ok']
    result = Presentation(copied)
    assert [s.shapes.title.text for s in result.slides] == ['B', 'A', 'A']
    for index in (1, 2):
        assert result.slides[index].shapes.title.click_action.target_slide == result.slides[0]
    assert validate_document(copied, strict=True)['issue_count'] == 0


def test_pptx_slide_id_mapping_is_checked_not_just_count(tmp_path):
    from zipfile import ZipFile
    from lxml import etree
    source, damaged = tmp_path / 'deck.pptx', tmp_path / 'damaged.pptx'
    prs = Presentation()
    for _ in range(2):
        prs.slides.add_slide(prs.slide_layouts[1])
    prs.save(source)
    with ZipFile(source) as archive, ZipFile(damaged, 'w') as target:
        for name in archive.namelist():
            data = archive.read(name)
            if name == 'ppt/presentation.xml':
                root = etree.fromstring(data)
                nodes = root.findall('.//{http://schemas.openxmlformats.org/presentationml/2006/main}sldId')
                attribute = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
                nodes[1].set(attribute, nodes[0].get(attribute))
                data = etree.tostring(root)
            target.writestr(name, data)
    assert 'SLIDE_RELATION_INVALID' in {i['code'] for i in validate_document(damaged)['issues']}


def test_pptx_custom_show_reference_blocks_deletion(tmp_path):
    from pptx.oxml.xmlchemy import OxmlElement as PptElement
    from pptx.oxml.ns import qn as ppt_qn
    source, output = tmp_path / 'deck.pptx', tmp_path / 'deleted.pptx'
    prs = Presentation()
    for _ in range(2):
        prs.slides.add_slide(prs.slide_layouts[1])
    shows = PptElement('p:custShowLst')
    show = PptElement('p:custShow'); show.set('name', 'seminar'); show.set('id', '1')
    slides = PptElement('p:sldLst')
    reference = PptElement('p:sld')
    reference.set(ppt_qn('r:id'), prs.slides._sldIdLst[1].rId)
    slides.append(reference); show.append(slides); shows.append(show)
    prs.part._element.append(shows)
    prs.save(source)
    result = run('pptx_slide', source, output, operation='delete', index=1)
    assert not result['ok'] and result['error']['code'] == 'E_FORMAT_UNSUPPORTED'
    assert 'presentation.xml' in result['error']['detail']
    assert not output.exists()
