"""真实文档边界：替换尾部、幻灯片关系、跨表范围和保存失败保护。"""
import json
from datetime import datetime
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation
from openpyxl import Workbook

from kiyori_office import protocol
from kiyori_office.paths import atomic_save, resolve_task_id
from kiyori_office.ops_docx.runs import compile_pattern, replace_in_paragraph


def test_cross_run_replace_preserves_suffix_and_later_matches():
    paragraph = Document().add_paragraph()
    paragraph.add_run('前文 AB').bold = True
    tail = paragraph.add_run('CD 后文 ABCD 尾部')
    tail.italic = True
    assert replace_in_paragraph(paragraph, compile_pattern('ABCD'), '新') == 2
    assert paragraph.text == '前文 新 后文 新 尾部'
    assert tail.italic is True


@pytest.mark.parametrize('index', [0, 1, 2])
def test_delete_slide_keeps_other_slides_and_relationships(tmp_path, index):
    source, output = tmp_path / 'source.pptx', tmp_path / 'out.pptx'
    deck = Presentation()
    for title in ('A', 'B', 'C'):
        deck.slides.add_slide(deck.slide_layouts[1]).shapes.title.text = title
    deck.save(source)
    result = protocol.run('pptx_slide', dict(path=str(source), output_path=str(output), operation='delete', index=index))
    assert result['ok'], result
    assert [s.shapes.title.text for s in Presentation(output).slides] == [t for i, t in enumerate(('A', 'B', 'C')) if i != index]
    assert protocol.run('office_validate', dict(path=str(output), strict=True))['ok']


def test_pptx_zero_layout_and_split_template_token(tmp_path):
    source = tmp_path / 'template.pptx'
    result = protocol.run('pptx_create', dict(slides=[{'title': 'Title'}], layout_index=0, output_path=str(source)))
    assert result['ok'], result
    deck = Presentation(source)
    assert deck.slides[0].slide_layout.name == deck.slide_layouts[0].name
    paragraph = deck.slides[0].shapes.title.text_frame.paragraphs[0]
    paragraph.clear()
    paragraph.add_run().text = '{{pro'
    tail = paragraph.add_run()
    tail.text = 'ject}} suffix'
    tail.font.bold = True
    deck.save(source)
    output = tmp_path / 'filled.pptx'
    result = protocol.run('pptx_template_fill', dict(path=str(source), output_path=str(output), variables={'project': 'Kiyori'}))
    assert result['ok'], result
    assert Presentation(output).slides[0].shapes.title.text == 'Kiyori suffix'


def test_xlsx_each_sheet_uses_own_range_and_serializes_dates(tmp_path):
    source = tmp_path / 'dates.xlsx'
    book = Workbook()
    book.active['A1'] = 'short'
    second = book.create_sheet('Long')
    second['C4'] = datetime(2026, 9, 10)
    book.save(source)
    result = protocol.run('xlsx_read', dict(path=str(source)))
    assert result['ok'], result
    encoded = json.dumps(result, ensure_ascii=False)
    assert '2026-09-10T00:00:00' in encoded
    sheets = result['data']['sheets']
    assert sheets[1]['range'] == 'A1:C4'


def test_atomic_save_failure_preserves_existing_file(tmp_path):
    target = tmp_path / 'valuable.docx'
    target.write_bytes(b'original')
    class FailingDocument:
        def save(self, path):
            Path(path).write_bytes(b'partial')
            raise OSError('disk full')
    with pytest.raises(OSError):
        atomic_save(FailingDocument(), target)
    assert target.read_bytes() == b'original'
    assert list(tmp_path.glob('.tmp-*')) == []


@pytest.mark.parametrize('task_id', ['.', '..', '../elsewhere'])
def test_task_id_rejects_parent_or_current_directory(task_id):
    with pytest.raises(protocol.OfficeError):
        resolve_task_id({'task_id': task_id})
