"""真实 OOXML 回读：混合写入、论文排版、PPT 布局与副本隔离。"""

from zipfile import ZipFile

import pytest

from kiyori_office import protocol


def call(command, **args):
    result = protocol.run(command, args)
    assert result['ok'], result
    return result


def test_rows_then_cells_and_final_counts(tmp_path):
    from openpyxl import load_workbook
    path = tmp_path / 'budget.xlsx'
    result = call('xlsx_write', rows=[['cost','revenue','profit'], [2,6,None]],
                  cells=[{'cell':'C2','formula':'=B2-A2'}, {'cell':'A1','value':None}], output_path=str(path))
    sheet = load_workbook(path).active
    assert sheet['C2'].value == '=B2-A2'
    assert sheet['A1'].value is None
    assert result['data']['formula_cells'] == 1
    assert result['data']['cells_written'] == 6
    assert result['data']['overwritten_in_batch'] == 2


@pytest.mark.parametrize('cells', [
    [{'cell':'A1','value':1,'formula':'=1'}], [{'cell':'A0','value':1}],
    [{'cell':'XFE1','value':1}], [{'cell':'A1:A2','value':1}],
    [{'cell':'A1048577','value':1}], [{'cell':'A1','formula':123}], []])
def test_invalid_batch_never_publishes(tmp_path, cells):
    path = tmp_path / 'invalid.xlsx'
    result = protocol.run('xlsx_write', {'cells':cells,'output_path':str(path)})
    assert not result['ok'] and result['error']['code'] == 'E_INPUT_SCHEMA', result
    assert not path.exists()


def test_strict_error_message_and_issues_include_location(tmp_path):
    path = tmp_path / 'formula.xlsx'
    call('xlsx_write', cells=[{'cell':'D2','formula':'=1+1'}], output_path=str(path))
    result = protocol.run('office_validate', {'path':str(path),'strict':True})
    assert not result['ok']
    assert 'FORMULA_NOT_RECALCULATED' in result['error']['message']
    assert 'D2' in result['error']['message']
    assert result['data']['issues'][0]['code'] == 'FORMULA_NOT_RECALCULATED'


def test_paper_layout_fonts_spacing_heading_and_table(tmp_path):
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.oxml.ns import qn
    path = tmp_path / 'paper.docx'
    call('docx_create', output_path=str(path), spec={
        'layout': {'page_setup':{'width_cm':21,'height_cm':29.7},
                   'margins_cm':{'top':2.54,'bottom':2.54,'left':3,'right':2.5},
                   'default_font':{'name':'Times New Roman','east_asia':'宋体','size_pt':12},
                   'paragraph_styles':{'Normal':{'line_spacing':1.5,'space_after_pt':0,'first_line_indent_cm':0.85,'widow_control':True},
                                       'Heading 1':{'east_asia':'黑体','size_pt':16,'bold':True,'keep_with_next':True,'page_break_before':True}}},
        'blocks':[{'type':'title','text':'论文题目'}, {'type':'heading','text':'研究方法','level':1},
                  {'type':'paragraph','runs':[{'text':'中文正文','east_asia':'宋体'}, {'text':' significant','bold':True}],
                   'format':{'alignment':'justify'}},
                  {'type':'table','rows':[['指标','值'], ['结果',None]],'column_widths_cm':[7,7]}]})
    doc = Document(path)
    assert abs(doc.sections[0].page_width - Cm(21)) < 1000
    normal = doc.styles['Normal']
    assert normal.font.name == 'Times New Roman'
    assert normal.element.rPr.rFonts.get(qn('w:eastAsia')) == '宋体'
    assert normal.paragraph_format.line_spacing == 1.5
    assert normal.paragraph_format.space_after == Pt(0)
    assert doc.styles['Heading 1'].paragraph_format.page_break_before
    assert doc.paragraphs[2].text == '中文正文 significant'
    assert doc.paragraphs[2].runs[-1].bold
    assert doc.tables[0].rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None
    assert doc.tables[0].cell(1,1).text == ''
    assert doc.tables[0].autofit is False
    call('office_validate',path=str(path),strict=True)


@pytest.mark.parametrize('layout', [
    {'paragraph_styles':{'Normal':{'line_spacing':1.5,'line_spacing_pt':20}}},
    {'paragraph_styles':{'Normal':{'alignment':'sideways'}}},
    {'paragraph_styles':{'Normal':{'size_pt':-1}}},
    {'margins_cm':{'left':20,'right':20}},
    {'default_font':{'east_asia':123}},
    {'paragraph_styles':{'Missing':{'bold':True}}}])
def test_bad_paper_layout_is_actionable_and_preserves_output(tmp_path, layout):
    path = tmp_path / 'existing.docx'
    path.write_bytes(b'existing')
    result = protocol.run('docx_create', {'spec':{'layout':layout,'blocks':[{'text':'x'}]},'output_path':str(path),'overwrite':True})
    assert not result['ok'] and result['error']['code'] == 'E_INPUT_SCHEMA', result
    assert path.read_bytes() == b'existing'


def test_docx_append_image_fits_final_section(tmp_path):
    from PIL import Image
    from docx import Document
    from docx.shared import Cm
    source = tmp_path / 'source.docx'
    document = Document()
    last = document.add_section()
    last.page_width = Cm(12)
    last.left_margin = last.right_margin = Cm(2)
    document.save(source)
    picture = tmp_path / 'portrait.png'
    Image.new('RGB',(200,1000),'red').save(picture)
    output = tmp_path / 'image.docx'
    call('docx_insert_image',path=str(source),image_path=str(picture),output_path=str(output))
    doc = Document(output)
    shape = doc.inline_shapes[0]
    assert shape.width <= Cm(8)
    assert shape.height <= doc.sections[-1].page_height - doc.sections[-1].top_margin - doc.sections[-1].bottom_margin


def test_page_numbers_are_real_fields_and_reapplying_does_not_duplicate(tmp_path):
    from docx import Document
    from docx.oxml.ns import qn
    source = tmp_path / 'paper.docx'
    call('docx_create',markdown='# 论文\n正文',output_path=str(source))
    output = tmp_path / 'numbered.docx'
    settings = {'header':{'text':'研究论文'},'footer':{'text':'第','page_number':True,'total_pages':True}}
    result = call('docx_style',path=str(source),header_footer=settings,output_path=str(output))
    assert result['warnings'][0]['code'] == 'FIELD_REFRESH_REQUIRED'
    call('docx_style',path=str(output),header_footer=settings,output_path=str(output),overwrite=True)
    doc = Document(output)
    assert doc.sections[0].header.paragraphs[0].text == '研究论文'
    fields = list(doc.sections[0].footer._element.iter(qn('w:instrText')))
    assert [f.text.strip() for f in fields] == ['PAGE','NUMPAGES']
    assert doc.settings.element.find(qn('w:updateFields')).get(qn('w:val')) == 'true'


@pytest.mark.parametrize('kind', ['column','bar','line','pie'])
def test_native_pptx_chart_roundtrip(tmp_path, kind):
    from pptx import Presentation
    path = tmp_path / (kind+'.pptx')
    call('pptx_create',output_path=str(path),slides=[{'layout_index':6,'elements':[
        {'type':'chart','left_cm':1,'top_cm':1,'width_cm':20,'height_cm':12,'chart_type':kind,
         'title':'季度销售','font_name':'Microsoft YaHei','size_pt':16,'categories':['第一季','第二季'],
         'series':[{'name':'销售额','values':[4,7]}]}]}])
    chart = Presentation(path).slides[0].shapes[0].chart
    assert list(chart.series[0].values) == [4,7]
    assert chart.chart_title.text_frame.text == '季度销售'
    call('office_validate',path=str(path),strict=True)


def test_pptx_set_text_preserves_designed_typography(tmp_path):
    from pptx import Presentation
    source = tmp_path / 'source.pptx'; output = tmp_path / 'edited.pptx'
    call('pptx_create',output_path=str(source),slides=[{'layout_index':6,'elements':[
        {'name':'title','left_cm':1,'top_cm':1,'width_cm':20,'height_cm':4,'text':'旧标题',
         'size_pt':32,'bold':True,'font_name':'Microsoft YaHei','color_rgb':'123456','alignment':'center'}]}])
    call('pptx_edit',path=str(source),slide_index=0,shape_name='title',text='新标题\n第二行',output_path=str(output))
    paragraphs = Presentation(output).slides[0].shapes[0].text_frame.paragraphs
    assert len(paragraphs) == 2
    for paragraph in paragraphs:
        run = paragraph.runs[0]
        assert run.font.size.pt == 32 and run.font.bold
        assert str(run.font.color.rgb) == '123456'


def test_format_reports_cache_invalidation(tmp_path):
    source = tmp_path / 'source.xlsx'; output = tmp_path / 'formatted.xlsx'
    call('xlsx_write',rows=[[1,'=A1+1']],output_path=str(source))
    result = call('xlsx_format',path=str(source),range='A1:B1',font={'bold':True},output_path=str(output))
    assert result['warnings'][0]['code'] == 'FORMULA_CACHE_INVALIDATED'
    assert 'xlsx_recalc' in result['next_actions']


def test_sheet_copy_does_not_silently_drop_charts(tmp_path):
    source = tmp_path / 'source.xlsx'; chart = tmp_path / 'chart.xlsx'; target = tmp_path / 'copied.xlsx'
    call('xlsx_write',rows=[['label','value'],['a',2]],output_path=str(source))
    call('xlsx_chart',path=str(source),data_range='A1:B2',output_path=str(chart))
    result = protocol.run('xlsx_sheet',{'path':str(chart),'operation':'copy','name':'Sheet','output_path':str(target)})
    assert not result['ok'] and result['error']['code'] == 'E_FORMAT_UNSUPPORTED'
    assert not target.exists()


def test_conversion_rejects_mislabeled_output_before_engine(tmp_path):
    source = tmp_path / 'source.docx'; output = tmp_path / 'wrong.xlsx'
    call('docx_create',markdown='sample',output_path=str(source))
    result = protocol.run('office_convert',{'from_path':str(source),'to_format':'pdf','engine':'libreoffice','output_path':str(output)})
    assert not result['ok'] and result['error']['code'] == 'E_INPUT_SCHEMA'
    assert not output.exists()


def test_scatter_chart_is_not_rejected_as_missing_category_axis(tmp_path):
    from pptx import Presentation
    from pptx.chart.data import XyChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    path = tmp_path / 'scatter.pptx'
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = XyChartData(); series = data.add_series('measured'); series.add_data_point(1,2); series.add_data_point(2,3)
    slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(1), Inches(1), Inches(4), Inches(3), data)
    prs.save(path)
    call('office_validate',path=str(path),strict=True)


def test_docx_style_in_place_changes_actual_source(tmp_path):
    from docx import Document
    source = tmp_path / 'source.docx'
    call('docx_create',markdown='原地排版',output_path=str(source))
    result = call('docx_style',path=str(source),default_font={'size_pt':18},in_place=True)
    assert result['artifacts'][0]['path'] == str(source.resolve())
    assert Document(source).styles['Normal'].font.size.pt == 18


def test_encrypted_pdf_reports_encryption_not_corruption(tmp_path):
    from pypdf import PdfWriter
    path = tmp_path / 'encrypted.pdf'
    writer = PdfWriter(); writer.add_blank_page(200,200); writer.encrypt('secret')
    with path.open('wb') as stream:
        writer.write(stream)
    result = call('office_validate',path=str(path),strict=False)
    assert result['data']['issues'][0]['code'] == 'PDF_ENCRYPTED'


def test_pptx_editable_layout_and_blank_title(tmp_path):
    from pptx import Presentation
    path = tmp_path / 'deck.pptx'
    call('pptx_create',output_path=str(path),slide_size_cm={'width':33.867,'height':19.05}, slides=[{
        'layout_index':6,'title':'研究结论','background_rgb':'F5F7FA','notes':'答辩讲稿',
        'elements':[{'type':'shape','name':'accent','shape':'rectangle','left_cm':1,'top_cm':5,'width_cm':0.2,'height_cm':8,'fill_rgb':'3366AA'},
                    {'type':'text','name':'findings','left_cm':2,'top_cm':5,'width_cm':15,'height_cm':8,
                     'font_name':'Microsoft YaHei','size_pt':24,'color_rgb':'223344',
                     'paragraphs':[{'text':'可编辑的中文结论','bullet':True}, {'text':'支持分层要点','bullet':True,'level':1}]},
                    {'type':'table','left_cm':19,'top_cm':5,'width_cm':12,'height_cm':6,'size_pt':18,
                     'rows':[['指标','结果'],['准确率','95%']]}]}])
    prs = Presentation(path)
    slide = prs.slides[0]
    assert any(s.has_text_frame and s.text == '研究结论' for s in slide.shapes)
    assert slide.notes_slide.notes_text_frame.text == '答辩讲稿'
    text = next(s for s in slide.shapes if s.name == 'findings')
    assert text.text_frame.paragraphs[1].level == 1
    assert text.text_frame.paragraphs[0].runs[0].font.size.pt == 24
    assert 'Microsoft YaHei' in text._element.xml
    assert any(s.has_table for s in slide.shapes)
    call('office_validate',path=str(path),strict=True)


def test_duplicate_notes_chart_and_background_are_independent(tmp_path):
    from pptx import Presentation
    from pptx.chart.data import ChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.dml.color import RGBColor
    from pptx.util import Inches
    source = tmp_path / 'chart.pptx'
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.notes_slide.notes_text_frame.text = 'original notes'
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(12,34,56)
    data = ChartData(); data.categories = ['A','B']; data.add_series('Sales',[1,2])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), data)
    prs.save(source)
    output = tmp_path / 'copy.pptx'
    call('pptx_slide',path=str(source),operation='duplicate',index=0,output_path=str(output))
    prs = Presentation(output)
    assert prs.slides[0].notes_slide.part is not prs.slides[1].notes_slide.part
    assert prs.slides[0].shapes[0].chart.part is not prs.slides[1].shapes[0].chart.part
    assert prs.slides[1].background.fill.fore_color.rgb == RGBColor(12,34,56)
    prs.slides[1].notes_slide.notes_text_frame.text = 'changed notes'
    replacement = ChartData(); replacement.categories = ['A','B']; replacement.add_series('Sales',[9,8])
    prs.slides[1].shapes[0].chart.replace_data(replacement)
    prs.save(output)
    check = Presentation(output)
    assert check.slides[0].notes_slide.notes_text_frame.text == 'original notes'
    assert list(check.slides[0].shapes[0].chart.series[0].values) == [1,2]
    assert list(check.slides[1].shapes[0].chart.series[0].values) == [9,8]
    with ZipFile(output) as archive:
        assert len(archive.namelist()) == len(set(archive.namelist()))
    call('office_validate',path=str(output),strict=True)


@pytest.mark.parametrize('dimensions', [{'width_emu':720000}, {'height_emu':1080000}, {}])
def test_pptx_image_one_dimension_respected_and_default_fits(tmp_path, dimensions):
    from PIL import Image
    from pptx import Presentation
    source = tmp_path / 'source.pptx'
    call('pptx_create',slides=[{'title':'portrait'}],output_path=str(source))
    picture = tmp_path / 'portrait.png'; Image.new('RGB',(200,1000),'red').save(picture)
    output = tmp_path / 'placed.pptx'
    call('pptx_media',path=str(source),slide_index=0,image_path=str(picture),output_path=str(output),**dimensions)
    prs = Presentation(output); image = list(prs.slides[0].shapes)[-1]
    for key, expected in dimensions.items():
        assert getattr(image,key[:-4]) == expected
    assert image.top + image.height <= prs.slide_height
    assert image.left + image.width <= prs.slide_width
