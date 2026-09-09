"""Word 的页面、字体和段落排版；创建与编辑共用一个实现。"""

import math

from ..protocol import OfficeError


def number(value, field, minimum=0, maximum=1000):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or not minimum <= value <= maximum:
        raise OfficeError("E_INPUT_SCHEMA", "%s 必须是 %s 到 %s 的有限数字" % (field, minimum, maximum))
    return value


def checked(config, keys, field):
    if not isinstance(config, dict):
        raise OfficeError("E_INPUT_SCHEMA", "%s 必须是对象" % field)
    extra = set(config) - set(keys)
    if extra:
        raise OfficeError("E_INPUT_SCHEMA", "%s 含不支持的字段：%s" % (field, ", ".join(sorted(extra))))
    return config


FONT_KEYS = {"name", "east_asia", "size_pt", "bold", "italic", "color_rgb"}
PARAGRAPH_KEYS = {"alignment", "line_spacing", "line_spacing_pt", "space_before_pt", "space_after_pt",
                  "first_line_indent_cm", "left_indent_cm", "right_indent_cm", "keep_with_next",
                  "keep_together", "page_break_before", "widow_control"}


def font_style(font, config):
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    import re

    for key in ("name", "east_asia"):
        if key in config and (not isinstance(config[key], str) or not config[key].strip()):
            raise OfficeError("E_INPUT_SCHEMA", "%s 必须是非空字体名称" % key)
    if "name" in config:
        font.name = config["name"]
    if "east_asia" in config:
        # font.name 只设置 ascii/hAnsi；不写 eastAsia，中文仍会被主题字体覆盖。
        fonts = font._element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn("w:eastAsia"), config["east_asia"])
        fonts.attrib.pop(qn("w:eastAsiaTheme"), None)
    if "size_pt" in config:
        font.size = Pt(number(config["size_pt"], "size_pt", 1, 400))
    for key in ("bold", "italic"):
        if key in config:
            if not isinstance(config[key], bool):
                raise OfficeError("E_INPUT_SCHEMA", "%s 必须是 boolean" % key)
            setattr(font, key, config[key])
    if "color_rgb" in config:
        color = config["color_rgb"]
        if not isinstance(color, str) or not re.fullmatch(r"[0-9a-fA-F]{6}", color):
            raise OfficeError("E_INPUT_SCHEMA", "color_rgb 必须是六位 RGB 十六进制字符串")
        font.color.rgb = RGBColor.from_string(color)


def paragraph_style(paragraph_format, config):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    if "alignment" in config:
        alignment = config["alignment"]
        values = {"LEFT": WD_ALIGN_PARAGRAPH.LEFT, "CENTER": WD_ALIGN_PARAGRAPH.CENTER,
                  "RIGHT": WD_ALIGN_PARAGRAPH.RIGHT, "JUSTIFY": WD_ALIGN_PARAGRAPH.JUSTIFY}
        if not isinstance(alignment, str) or alignment.upper() not in values:
            raise OfficeError("E_INPUT_SCHEMA", "alignment 必须是 left/center/right/justify")
        paragraph_format.alignment = values[alignment.upper()]
    if "line_spacing" in config and "line_spacing_pt" in config:
        raise OfficeError("E_INPUT_SCHEMA", "倍数行距 line_spacing 与固定行距 line_spacing_pt 不能同时指定")
    if "line_spacing" in config:
        paragraph_format.line_spacing = number(config["line_spacing"], "line_spacing", 0.5, 10)
    if "line_spacing_pt" in config:
        paragraph_format.line_spacing = Pt(number(config["line_spacing_pt"], "line_spacing_pt", 1, 400))
    for key in ("space_before_pt", "space_after_pt"):
        if key in config:
            setattr(paragraph_format, key[:-3], Pt(number(config[key], key, 0, 400)))
    for key in ("first_line_indent_cm", "left_indent_cm", "right_indent_cm"):
        if key in config:
            setattr(paragraph_format, key[:-3], Cm(number(config[key], key, -10, 30)))
    for key in ("keep_with_next", "keep_together", "page_break_before", "widow_control"):
        if key in config:
            if not isinstance(config[key], bool):
                raise OfficeError("E_INPUT_SCHEMA", "%s 必须是 boolean" % key)
            setattr(paragraph_format, key, config[key])


def apply_layout(document, config):
    from docx.shared import Cm
    from docx.enum.style import WD_STYLE_TYPE

    checked(config, {"margins_cm", "default_font", "paragraph_styles", "page_setup", "header_footer"}, "layout")
    applied = []
    if "default_font" in config:
        font_style(document.styles["Normal"].font, checked(config["default_font"], FONT_KEYS, "default_font"))
        applied.append("default_font")
    if "paragraph_styles" in config:
        styles = config["paragraph_styles"]
        if not isinstance(styles, dict):
            raise OfficeError("E_INPUT_SCHEMA", "paragraph_styles 必须是对象")
        for name, values in styles.items():
            checked(values, FONT_KEYS | PARAGRAPH_KEYS, "paragraph_styles.%s" % name)
            if name not in document.styles or document.styles[name].type != WD_STYLE_TYPE.PARAGRAPH:
                raise OfficeError("E_INPUT_SCHEMA", "未知段落样式：%s" % name)
            style = document.styles[name]
            font_style(style.font, values)
            paragraph_style(style.paragraph_format, values)
            applied.append("style:%s" % name)
    margins = checked(config.get("margins_cm", {}), {"top", "bottom", "left", "right"}, "margins_cm")
    page = checked(config.get("page_setup", {}), {"width_cm", "height_cm", "header_distance_cm", "footer_distance_cm"}, "page_setup")
    for section in document.sections:
        for key, value in margins.items():
            setattr(section, key + "_margin", Cm(number(value, "margins_cm." + key, 0, 50)))
        for key, value in page.items():
            attr = {"width_cm": "page_width", "height_cm": "page_height"}.get(key, key[:-3])
            setattr(section, attr, Cm(number(value, "page_setup." + key, 1 if key in ("width_cm", "height_cm") else 0, 55.88)))
        if section.page_width <= section.left_margin + section.right_margin or section.page_height <= section.top_margin + section.bottom_margin:
            raise OfficeError("E_INPUT_SCHEMA", "页边距占满页面，正文区域必须大于零")
    if margins:
        applied.append("margins")
    if page:
        applied.append("page_setup")
    if "header_footer" in config:
        header_footer(document, config["header_footer"])
        applied.append("header_footer")
    return applied


def append_field(paragraph, instruction, placeholder):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    code = OxmlElement("w:instrText"); code.set(qn("xml:space"), "preserve"); code.text = " " + instruction + " "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t"); value.text = placeholder
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for item in (begin, code, separate, value, end):
        run._r.append(item)


def header_footer(document, config):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    checked(config, {"header", "footer"}, "header_footer")
    seen = set()
    for section in document.sections:
        for region, value in config.items():
            checked(value, {"text", "alignment", "page_number", "total_pages"} if region == "footer" else {"text", "alignment"}, region)
            if not isinstance(value.get("text", ""), str):
                raise OfficeError("E_INPUT_SCHEMA", "%s.text 必须为字符串" % region)
            story = getattr(section, region)
            element = story._element
            if element in seen:
                continue
            seen.add(element)
            # header/footer 参数明确表示替换该默认区域，不修改首页或偶数页区域。
            for child in list(element):
                element.remove(child)
            paragraph = story.add_paragraph(value.get("text", ""), style=region.title())
            paragraph_style(paragraph.paragraph_format, {"alignment": value.get("alignment", "center"), "first_line_indent_cm": 0})
            for field, instruction in (("page_number", "PAGE"), ("total_pages", "NUMPAGES")):
                if field in value and not isinstance(value[field], bool):
                    raise OfficeError("E_INPUT_SCHEMA", "%s 必须为 boolean" % field)
                if value.get(field):
                    if paragraph.text:
                        paragraph.add_run(" / " if field == "total_pages" else " ")
                    append_field(paragraph, instruction, "1")
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields"); settings.append(update)
    update.set(qn("w:val"), "true")


def layout_warnings(config):
    footer = config.get("header_footer", {}).get("footer", {})
    if footer.get("page_number") or footer.get("total_pages"):
        return [{"code": "FIELD_REFRESH_REQUIRED", "message": "已插入真实页码域；当前缓存是占位值，必须在 Word/LibreOffice 更新域并渲染检查，不能按占位页码交付。"}]
    return []


def format_table(table, rows, *, header=True, widths=None):
    from docx.oxml import OxmlElement
    from docx.shared import Cm

    if header:
        properties = table.rows[0]._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:tblHeader"))
        for cell in table.rows[0].cells:
            for run in cell.paragraphs[0].runs:
                run.bold = True
    if widths is not None:
        if not isinstance(widths, list) or len(widths) != len(table.columns):
            raise OfficeError("E_INPUT_SCHEMA", "column_widths_cm 长度必须与列数一致")
        table.autofit = False
        for col, value in enumerate(widths):
            width = Cm(number(value, "column_widths_cm", 0.1, 50))
            table.columns[col].width = width
            for row in table.rows:
                row.cells[col].width = width


def table_rows(rows):
    if not isinstance(rows, list) or not rows or not all(isinstance(row, list) for row in rows) or not max(map(len, rows), default=0):
        raise OfficeError("E_INPUT_SCHEMA", "表格 rows 必须是至少含一个单元格的二维数组")
    if len(rows) * max(map(len, rows)) > 100000:
        raise OfficeError("E_BUDGET_EXCEEDED", "表格超过 100000 单元格")
    return max(map(len, rows))
