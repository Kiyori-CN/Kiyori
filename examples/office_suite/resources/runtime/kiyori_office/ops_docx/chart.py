"""复用原生图表生成器，把图表及独立嵌入工作簿接入 DOCX OPC 包。"""

from ..protocol import OfficeError


def copy_chart_for_merge(source, package, copied):
    """图表与可变工作簿使用独立 OPC 部件；不能复用源包名称或直接复制 rId。"""
    from pathlib import PurePosixPath
    from docx.opc.part import Part
    from docx.opc.packuri import PackURI
    from lxml import etree

    if source in copied:
        return copied[source]
    used = {str(part.partname) for part in package.iter_parts()}
    used.update(str(part.partname) for part in copied.values())
    original = PurePosixPath(str(source.partname))
    index = 1
    name = str(original)
    while name in used:
        name = str(original.with_name("%s_merge%d%s" % (original.stem, index, original.suffix)))
        index += 1
    target = Part(PackURI(name), source.content_type, source.blob, package)
    copied[source] = target
    mapping = {}
    for relation in source.rels.values():
        kind = relation.reltype.rsplit('/', 1)[-1]
        if kind not in ('package', 'image', 'chartStyle', 'chartColorStyle'):
            raise OfficeError('E_FORMAT_UNSUPPORTED', '图表包含未支持的关系类型', detail=relation.reltype)
        if relation.is_external:
            mapping[relation.rId] = target.relate_to(relation.target_ref, relation.reltype, is_external=True)
        else:
            child = copy_chart_for_merge(relation.target_part, package, copied)
            mapping[relation.rId] = target.relate_to(child, relation.reltype)
    if mapping:
        root = etree.fromstring(source.blob)
        for node in root.iter():
            for key, value in list(node.attrib.items()):
                if key.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):
                    if value not in mapping:
                        raise OfficeError('E_FORMAT_UNSUPPORTED', '图表关系引用不存在', detail=value)
                    node.set(key, mapping[value])
        target._blob = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
    return target


def append_chart(document, block, section=None):
    from pptx import Presentation
    from docx.oxml import parse_xml
    from docx.opc.part import Part
    from docx.opc.packuri import PackURI
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from .layout import checked, number, paragraph_style
    from ..ops_pptx.layout import add_elements
    from lxml import etree
    checked(block, {"type", "name", "width_cm", "height_cm", "categories", "series", "chart_type", "title", "chart_style", "caption"}, "chart")
    section = section or document.sections[-1]
    width = number(block.get("width_cm", 14), "width_cm", 1, 50)
    height = number(block.get("height_cm", 8), "height_cm", 1, 50)
    if width*360000 > section.page_width-section.left_margin-section.right_margin or height*360000 > section.page_height-section.top_margin-section.bottom_margin:
        raise OfficeError("E_INPUT_SCHEMA", "图表尺寸超出正文区域")
    prs = Presentation()
    prs.slide_width = int(width*360000)
    prs.slide_height = int(height*360000)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    spec = {k:v for k,v in block.items() if k not in {"caption"}}
    spec.update(type="chart",left_cm=0,top_cm=0,width_cm=width,height_cm=height)
    chart_shape = add_elements(prs,slide,[spec])[0]
    package = document.part.package
    used = {str(part.partname) for part in package.iter_parts()}
    copied = {}
    def copy_part(part):
        if part in copied:
            return copied[part]
        if part.content_type.endswith("chart+xml"):
            template = "/word/charts/chart%d.xml"
        elif part.content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
            template = "/word/embeddings/chartData%d.xlsx"
        else:
            raise OfficeError("E_FORMAT_UNSUPPORTED", "图表生成器产生了未支持的关系类型")
        index=1
        while template % index in used:
            index+=1
        name=template % index
        used.add(name)
        target=Part(PackURI(name),part.content_type,part.blob,package)
        copied[part]=target
        mapping={}
        for relation in part.rels.values():
            if relation.is_external:
                raise OfficeError("E_FORMAT_UNSUPPORTED", "新建图表不能包含外部关系")
            mapping[relation.rId]=target.relate_to(copy_part(relation.target_part),relation.reltype)
        if mapping:
            root=etree.fromstring(part.blob)
            for node in root.iter():
                for key,value in list(node.attrib.items()):
                    if key.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}") and value in mapping:
                        node.set(key,mapping[value])
            target._blob=etree.tostring(root,xml_declaration=True,encoding="UTF-8",standalone=True)
        return target
    relationship=document.part.relate_to(copy_part(chart_shape.chart.part),RT.CHART)
    identity=document.part.next_id
    xml='''<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
        xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
        xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"
        xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
        <wp:inline><wp:extent cx="%d" cy="%d"/><wp:docPr id="%d" name="Chart %d"/>
        <a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/chart">
        <c:chart r:id="%s"/></a:graphicData></a:graphic></wp:inline></w:drawing>''' % (width*360000,height*360000,identity,identity,relationship)
    paragraph=document.add_paragraph()
    paragraph_style(paragraph.paragraph_format,{"alignment":"center","first_line_indent_cm":0})
    paragraph.add_run()._r.append(parse_xml(xml))
    if block.get("caption"):
        paragraph.paragraph_format.keep_with_next=True
        document.add_paragraph(block["caption"],style="Caption")


def edit_chart(document, chart_node, config):
    """仅为选中的引用生成独立 chart/workbook，避免共享部件发生串改。"""
    from docx.opc.part import Part
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml.ns import qn
    from pptx import Presentation
    from pptx.parts.chart import ChartPart
    from pptx.opc.package import Part as PptPart
    from pptx.opc.packuri import PackURI
    from ..ops_pptx.style import chart_style, replace_chart_data
    from lxml import etree
    identity=chart_node.get(qn('r:id'))
    source=document.part.related_parts[identity]
    root=etree.fromstring(source.blob)
    external=root.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}externalData')
    workbook_identity=external.get(qn('r:id')) if external is not None else None
    if not workbook_identity or source.rels[workbook_identity].is_external:
        raise OfficeError('E_FORMAT_UNSUPPORTED','图表必须具有内部嵌入工作簿；不改外部链接图表')
    workbook=source.rels[workbook_identity].target_part
    if workbook.content_type != 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
        raise OfficeError('E_FORMAT_UNSUPPORTED','只编辑标准 XLSX 数据源图表')
    package=Presentation().part.package
    working=ChartPart.load(PackURI('/ppt/charts/chart1.xml'),source.content_type,package,source.blob)
    copied_workbook=PptPart.load(PackURI('/ppt/embeddings/data.xlsx'),workbook.content_type,package,workbook.blob)
    rid=working.relate_to(copied_workbook,RT.PACKAGE)
    working._element.xpath('./c:externalData')[0].set(qn('r:id'),rid)
    if 'chart_data' in config:
        replace_chart_data(working.chart,config['chart_data'])
    if 'chart_style' in config:
        chart_style(working.chart,config['chart_style'])
    root=etree.fromstring(working.blob)
    root.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}externalData').set(qn('r:id'),workbook_identity)
    destination=Part(document.part.package.next_partname('/word/charts/chart%d.xml'),source.content_type,
                     etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True),document.part.package)
    # 保留所有现存样式关系；仅数据源复制，重新映射关系身份。
    mapping={}
    for old_id,relation in source.rels.items():
        if relation.is_external:
            new_id=destination.relate_to(relation.target_ref,relation.reltype,is_external=True)
        else:
            target=relation.target_part
            if old_id==workbook_identity:
                target=Part(document.part.package.next_partname('/word/embeddings/chartData%d.xlsx'),workbook.content_type,copied_workbook.blob,document.part.package)
            new_id=destination.relate_to(target,relation.reltype)
        mapping[old_id]=new_id
    root=etree.fromstring(destination.blob)
    for node in root.iter():
        for key,value in list(node.attrib.items()):
            if key.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}') and value in mapping:
                node.set(key,mapping[value])
    destination._blob=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
    chart_node.set(qn('r:id'),document.part.relate_to(destination,RT.CHART))
    # 原引用在正文中不再出现时，断开旧部件；其他引用继续拥有原图表。
    if not any(node.get(qn('r:id'))==identity for node in document.element.iter()):
        document.part.drop_rel(identity)
