"""docx_from_template：{{变量}} 模板填充。

连续普通文字跨 run 单次匹配，保留未命中结构并避免替换值递归展开。
"""

from __future__ import annotations

import re
from typing import Any, Dict, List

from ..paths import atomic_save, artifact, resolve_output_path, resolve_path, resolve_task_id, task_dir
from ..protocol import OfficeError, engine_version, register
from ..readers.docx_reader import require_docx
from .runs import document_text_stories, text_run_groups, _replace_runs

VARIABLE_PATTERN = re.compile(r"\{\{\s*([A-Za-z0-9_.\-\u4e00-\u9fff]+)\s*\}\}")

@register(
    "docx_from_template",
    schema="docx_from_template",
    engine="python-docx",
    next_actions=["office_validate", "office_render_preview"],
)
def docx_from_template(args: Dict[str, Any]) -> Dict[str, Any]:
    require_docx()
    import docx  # type: ignore

    task_id = resolve_task_id(args)
    directory = task_dir(task_id)
    source = resolve_path(args.get("path"), args=args, field="path", must_exist=True)
    variables = args.get("variables")
    if not isinstance(variables, dict):
        raise OfficeError("E_INPUT_SCHEMA", "variables 必须是对象")
    strict = bool(args.get("strict", True))
    output = resolve_output_path(
        args.get("output_path"),
        args=args,
        field="output_path",
        default_dir=directory / "out",
        default_name=source.name,
        overwrite=bool(args.get("overwrite")),
    )

    document = docx.Document(str(source))
    used: List[str] = []
    missing: List[str] = []
    replacements = 0
    protected = []

    def value_for(match):
        name = match.group(1)
        if name not in variables:
            if name not in missing:
                missing.append(name)
            return None
        if name not in used:
            used.append(name)
        return str(variables[name])

    for story_index, story in enumerate(document_text_stories(document, include_headers=True)):
        state = [0]
        for paragraph_index, paragraph in enumerate(story):
            groups = text_run_groups(paragraph, state)
            available = sum(len(VARIABLE_PATTERN.findall(''.join(run.text for run in group))) for group in groups)
            if len(VARIABLE_PATTERN.findall(paragraph.text)) > available:
                protected.append({'story_index': story_index, 'xml_paragraph_index': paragraph_index})
            for group in groups:
                # 同一组仅扫描一次原始文本，从后向前按位置填充；值中的
                # {{...}} 是用户数据，不能被后续变量意外递归展开。
                replacements += _replace_runs(group, VARIABLE_PATTERN, value_for)

    if protected and strict:
        raise OfficeError('E_FORMAT_UNSUPPORTED', '模板标记位于域、链接或跨结构边界，未写入文件',
                          detail='protected_placeholders=%s' % protected,
                          data={'protected_placeholders': protected},
                          remedy='将模板标记放入连续普通文字；strict=false 仅保留无法处理的标记，并不代表填充完整。')

    if missing and strict:
        raise OfficeError(
            "E_TEMPLATE_VAR_MISSING",
            "模板变量缺少取值",
            detail="missing=%s" % ", ".join(missing),
            remedy="补齐 variables，或显式设置 strict=false 保留原样",
        )

    output.parent.mkdir(parents=True, exist_ok=True)
    atomic_save(document, output)
    return {
        "artifacts": [artifact(output)],
        "warnings": ([{'code': 'PROTECTED_TEMPLATE_MARKERS', 'message': '部分模板标记位于受保护结构内，已保留；见 data.protected_placeholders。'}] if protected else []),
        "data": {
            "used_variables": used,
            "missing_variables": missing,
            "replacements": replacements,
            "protected_placeholders": protected,
            "work_dir": str(directory),
        },
        "engine_version": engine_version("python-docx"),
    }
