"""docx_create / docx_table / docx_style。"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Dict, List, Optional

from ..budget import bounded_text
from ..paths import artifact, atomic_write_text, resolve_output_path, resolve_task_id, task_dir
from ..protocol import OfficeError, engine_version, register
from ..readers.docx_reader import docx_outline, require_docx

@register(
    "docx_create",
    schema="docx_create",
    engine="python-docx",
    next_actions=["docx_outline", "office_validate", "office_render_preview"],
)
def docx_create(args: Dict[str, Any]) -> Dict[str, Any]:
    require_docx()
    import docx  # type: ignore
    from docx.shared import Pt  # type: ignore

    task_id = resolve_task_id(args)
    directory = task_dir(task_id)
    output = resolve_output_path(
        args.get("output_path"),
        args=args,
        field="output_path",
        default_dir=directory / "out",
        default_name=str(args.get("file_name") or "document.docx"),
        overwrite=bool(args.get("overwrite")),
    )

    spec = args.get("spec")
    markdown = args.get("markdown")
    if not spec and not markdown:
        raise OfficeError(
            "E_INPUT_SCHEMA",
            "docx_create 需要 spec 或 markdown 之一",
            remedy="提供结构化 spec（blocks 数组）或 markdown 文本",
        )

    document = docx.Document()
    if markdown:
        _append_markdown(document, str(markdown), Pt)
    else:
        _append_spec(document, spec, Pt)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    outline = docx_outline(output)
    return {
        "artifacts": [artifact(output)],
        "data": {
            "work_dir": str(directory),
            "paragraph_count": outline["paragraph_count"],
            "table_count": outline["table_count"],
            "headings": outline["headings"],
        },
        "engine_version": engine_version("python-docx"),
    }


def _append_spec(document, spec: Any, pt) -> None:
    if not isinstance(spec, dict) or not isinstance(spec.get("blocks"), list):
        raise OfficeError(
            "E_INPUT_SCHEMA",
            "spec 必须是含 blocks 数组的对象",
            detail="spec=%r" % (type(spec).__name__,),
        )
    for index, block in enumerate(spec["blocks"]):
        if not isinstance(block, dict):
            raise OfficeError("E_INPUT_SCHEMA", "spec.blocks[%d] 必须是对象" % index)
        kind = str(block.get("type") or "paragraph")
        if kind in ("heading", "title"):
            level = int(block.get("level") or (0 if kind == "title" else 1))
            text = str(block.get("text") or "")
            if level <= 0:
                document.add_heading(text, level=0)
            else:
                document.add_heading(text, level=min(level, 9))
        elif kind == "paragraph":
            paragraph = document.add_paragraph(str(block.get("text") or ""))
            style = block.get("style")
            if style:
                try:
                    paragraph.style = document.styles[str(style)]
                except KeyError as exc:
                    raise OfficeError(
                        "E_INPUT_SCHEMA",
                        "未知段落样式: %s" % style,
                        detail=str(exc),
                    ) from exc
        elif kind == "bullet":
            document.add_paragraph(str(block.get("text") or ""), style="List Bullet")
        elif kind == "number":
            document.add_paragraph(str(block.get("text") or ""), style="List Number")
        elif kind == "table":
            rows = block.get("rows")
            if not isinstance(rows, list) or not rows:
                raise OfficeError("E_INPUT_SCHEMA", "table 需要非空 rows 数组")
            columns = max(len(row) for row in rows if isinstance(row, list))
            table = document.add_table(rows=len(rows), cols=columns)
            table.style = str(block.get("style") or "Table Grid")
            for row_index, row in enumerate(rows):
                if not isinstance(row, list):
                    raise OfficeError("E_INPUT_SCHEMA", "table.rows[%d] 必须是数组" % row_index)
                for column_index, value in enumerate(row):
                    table.cell(row_index, column_index).text = str(value)
        elif kind == "page_break":
            document.add_page_break()
        else:
            raise OfficeError(
                "E_INPUT_SCHEMA",
                "不支持的 spec 块类型: %s" % kind,
                remedy="使用 heading/title/paragraph/bullet/number/table/page_break",
            )


_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_MD_BULLET = re.compile(r"^[-*+]\s+(.*)$")
_MD_NUMBER = re.compile(r"^\d+[.)]\s+(.*)$")


def _append_markdown(document, markdown: str, pt) -> None:
    lines = markdown.replace("\r\n", "\n").split("\n")
    table_buffer: List[List[str]] = []

    def flush_table() -> None:
        nonlocal table_buffer
        if not table_buffer:
            return
        columns = max(len(row) for row in table_buffer)
        table = document.add_table(rows=len(table_buffer), cols=columns)
        table.style = "Table Grid"
        for row_index, row in enumerate(table_buffer):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = value
        table_buffer = []

    for raw_line in lines:
        line = raw_line.rstrip()
        if line.strip().startswith("|") and line.strip().endswith("|"):
            cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                continue
            table_buffer.append(cells)
            continue
        flush_table()
        if not line.strip():
            continue
        heading = _MD_HEADING.match(line)
        if heading:
            document.add_heading(heading.group(2).strip(), level=len(heading.group(1)))
            continue
        bullet = _MD_BULLET.match(line)
        if bullet:
            document.add_paragraph(bullet.group(1).strip(), style="List Bullet")
            continue
        numbered = _MD_NUMBER.match(line)
        if numbered:
            document.add_paragraph(numbered.group(1).strip(), style="List Number")
            continue
        document.add_paragraph(line)
    flush_table()


@register(
    "docx_table",
    schema="docx_table",
    engine="python-docx",
    next_actions=["docx_outline", "office_validate"],
)
def docx_table(args: Dict[str, Any]) -> Dict[str, Any]:
    require_docx()
    import docx  # type: ignore
    from docx.shared import Cm  # type: ignore

    task_id = resolve_task_id(args)
    directory = task_dir(task_id)
    source = Path(args["path"]).resolve()
    if not source.is_file():
        raise OfficeError("E_PATH_INVALID", "输入文件不存在", detail=str(source))
    output = resolve_output_path(
        args.get("output_path"),
        args=args,
        field="output_path",
        default_dir=directory / "out",
        default_name=source.name,
        overwrite=bool(args.get("overwrite")),
        in_place=bool(args.get("in_place")),
    )
    document = docx.Document(str(source))

    rows = args.get("rows")
    if not isinstance(rows, list) or not rows:
        raise OfficeError("E_INPUT_SCHEMA", "rows 必须是非空二维数组")
    header = bool(args.get("header", True))
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.style = str(args.get("style") or "Table Grid")
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = str(value)
    widths = args.get("column_widths_cm")
    if widths:
        if len(widths) != columns:
            raise OfficeError(
                "E_INPUT_SCHEMA",
                "column_widths_cm 长度必须与列数一致",
                detail="columns=%d widths=%d" % (columns, len(widths)),
            )
        # 列宽与单元格宽度必须同单位同时设置，否则 Word 只认其中一处
        for column_index, width in enumerate(widths):
            for row in table.rows:
                row.cells[column_index].width = Cm(float(width))
            table.columns[column_index].width = Cm(float(width))
    document.save(str(output))
    return {
        "artifacts": [artifact(output)],
        "data": {
            "rows": len(rows),
            "columns": columns,
            "header": header,
            "work_dir": str(directory),
        },
        "engine_version": engine_version("python-docx"),
    }


@register(
    "docx_style",
    schema="docx_style",
    engine="python-docx",
    next_actions=["office_validate", "office_render_preview"],
)
def docx_style(args: Dict[str, Any]) -> Dict[str, Any]:
    require_docx()
    import docx  # type: ignore
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    from docx.shared import Cm, Pt  # type: ignore

    task_id = resolve_task_id(args)
    directory = task_dir(task_id)
    source = Path(args["path"]).resolve()
    if not source.is_file():
        raise OfficeError("E_PATH_INVALID", "输入文件不存在", detail=str(source))
    output = resolve_output_path(
        args.get("output_path"),
        args=args,
        field="output_path",
        default_dir=directory / "out",
        default_name=source.name,
        overwrite=bool(args.get("overwrite")),
        in_place=bool(args.get("in_place")),
    )
    document = docx.Document(str(source))
    applied: List[str] = []

    margins = args.get("margins_cm")
    if isinstance(margins, dict):
        for section in document.sections:
            if "top" in margins:
                section.top_margin = Cm(float(margins["top"]))
            if "bottom" in margins:
                section.bottom_margin = Cm(float(margins["bottom"]))
            if "left" in margins:
                section.left_margin = Cm(float(margins["left"]))
            if "right" in margins:
                section.right_margin = Cm(float(margins["right"]))
        applied.append("margins")

    default_font = args.get("default_font")
    if isinstance(default_font, dict):
        style = document.styles["Normal"]
        if default_font.get("name"):
            style.font.name = str(default_font["name"])
        if default_font.get("size_pt"):
            style.font.size = Pt(float(default_font["size_pt"]))
        applied.append("default_font")

    paragraph_styles = args.get("paragraph_styles")
    if isinstance(paragraph_styles, dict):
        for name, config in paragraph_styles.items():
            if not isinstance(config, dict):
                continue
            try:
                style = document.styles[str(name)]
            except KeyError as exc:
                raise OfficeError(
                    "E_INPUT_SCHEMA", "未知样式: %s" % name, detail=str(exc)
                ) from exc
            if config.get("size_pt"):
                style.font.size = Pt(float(config["size_pt"]))
            if config.get("bold") is not None:
                style.font.bold = bool(config["bold"])
            if config.get("color_rgb"):
                from docx.shared import RGBColor  # type: ignore

                style.font.color.rgb = RGBColor.from_string(str(config["color_rgb"]))
            if config.get("alignment"):
                style.paragraph_format.alignment = getattr(
                    WD_ALIGN_PARAGRAPH, str(config["alignment"]).upper()
                )
            applied.append("style:%s" % name)

    document.save(str(output))
    return {
        "artifacts": [artifact(output)],
        "data": {"applied": applied, "work_dir": str(directory)},
        "engine_version": engine_version("python-docx"),
    }
