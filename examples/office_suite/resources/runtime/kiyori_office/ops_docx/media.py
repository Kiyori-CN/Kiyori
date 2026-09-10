"""docx_insert_image / docx_extract_media。"""

from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any, Dict, List

from ..paths import (
    atomic_save,
    artifact,
    atomic_write_bytes,
    resolve_output_dir,
    resolve_output_path,
    resolve_path,
    resolve_task_id,
    sanitize_filename,
    task_dir,
)
from ..protocol import OfficeError, engine_version, register
from ..readers.docx_reader import require_docx


def append_image(document, image, config, section=None):
    from .layout import number, paragraph_style
    from docx.shared import Cm, Emu
    from PIL import Image
    section = section or document.sections[-1]
    available_width = section.page_width - section.left_margin - section.right_margin
    available_height = section.page_height - section.top_margin - section.bottom_margin
    with Image.open(image) as picture:
        ratio = picture.height / picture.width
    if "width_cm" in config:
        width = Cm(number(config["width_cm"], "width_cm", 0.01, 100))
        if width > available_width or width * ratio > available_height:
            raise OfficeError("E_INPUT_SCHEMA", "图片显式尺寸超出正文区域，请减小 width_cm")
    else:
        width = min(available_width, int(available_height / ratio))
    paragraph = document.add_paragraph()
    paragraph_style(paragraph.paragraph_format, {"alignment": config.get("alignment", "center"), "first_line_indent_cm": 0})
    picture = paragraph.add_run().add_picture(str(image), width=Emu(int(width)))
    if "alt_text" in config:
        if not isinstance(config["alt_text"], str):
            raise OfficeError("E_INPUT_SCHEMA", "alt_text 必须是字符串")
        picture._inline.docPr.set("descr", config["alt_text"])
    if config.get("caption"):
        paragraph.paragraph_format.keep_with_next = True
        caption = document.add_paragraph(config["caption"], style="Caption")
        paragraph_style(caption.paragraph_format, {"alignment": "center", "first_line_indent_cm": 0})
    return int(width)

@register(
    "docx_insert_image",
    schema="docx_insert_image",
    engine="python-docx",
    next_actions=["office_validate", "office_render_preview"],
)
def docx_insert_image(args: Dict[str, Any]) -> Dict[str, Any]:
    require_docx()
    import docx  # type: ignore
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
    from docx.shared import Cm  # type: ignore

    task_id = resolve_task_id(args)
    directory = task_dir(task_id)
    source = resolve_path(args.get("path"), args=args, field="path", must_exist=True)
    image = resolve_path(args.get("image_path"), args=args, field="image_path", must_exist=True)
    output = resolve_output_path(
        args.get("output_path") or (str(source) if args.get("in_place") else None),
        args=args,
        field="output_path",
        default_dir=directory / "out",
        default_name=source.name,
        overwrite=bool(args.get("overwrite")),
        in_place=bool(args.get("in_place")),
    )
    document = docx.Document(str(source))
    width = append_image(document, image, args)
    output.parent.mkdir(parents=True, exist_ok=True)
    atomic_save(document, output)
    return {
        "artifacts": [artifact(output)],
        "data": {"image": str(image), "width_emu": int(width), "work_dir": str(directory)},
        "engine_version": engine_version("python-docx"),
    }


@register(
    "docx_extract_media",
    schema="docx_extract_media",
    engine="zipfile",
    next_actions=["office_read"],
)
def docx_extract_media(args: Dict[str, Any]) -> Dict[str, Any]:
    task_id = resolve_task_id(args)
    directory = task_dir(task_id)
    source = resolve_path(args.get("path"), args=args, field="path", must_exist=True)
    target_dir = resolve_output_dir(
        args.get("output_path"),
        args=args,
        field="output_path",
        default_dir=directory
        / "out"
        / sanitize_filename(str(args.get("target_dir_name") or (source.stem + "-media"))),
        overwrite=bool(args.get("overwrite")),
    )
    extracted: List[Dict[str, Any]] = []
    try:
        with zipfile.ZipFile(source) as archive:
            for name in archive.namelist():
                if not name.startswith("word/media/") or name.endswith("/"):
                    continue
                payload = archive.read(name)
                destination = target_dir / sanitize_filename(Path(name).name)
                atomic_write_bytes(destination, payload)
                extracted.append(artifact(destination))
    except zipfile.BadZipFile as exc:
        raise OfficeError(
            "E_DOC_CORRUPT", "DOCX 不是有效的 ZIP 容器", detail=str(exc)
        ) from exc
    return {
        "artifacts": extracted,
        "data": {
            "count": len(extracted),
            "target_dir": str(target_dir),
            "work_dir": str(directory),
        },
    }
